"""One protected, streaming ZIP from one checked customer result stream.

Run on the main thread of the dedicated export process (never the GUI process).
openpyxl's temporary XML is confined to the protected run directory. No workbook
or customer population is accumulated in RAM; product aggregates are catalog-sized.
"""

from collections import Counter
import csv
from datetime import datetime, timezone, timedelta
from decimal import Decimal
import math
import os
from pathlib import Path
import re
import shutil
import sqlite3
import tempfile
from time import perf_counter
import zipfile

from openpyxl import Workbook, load_workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.writer.excel import ExcelWriter

from Application.customer_contact_search import ContactType, normalize_search_input
from Application.recommendation_export_records import (
    ExportError,
    ExportCancelled,
    ExportSummary,
)

ONLINE = "RecommendationsOnlineStore.csv"
MINDBOX = "RecommendationsMindbox.csv"
ANALYTICS = "RecommendationsAnalytics.xlsx"
MEMBERS = (ONLINE, MINDBOX, ANALYTICS)
ONLINE_HEADERS = ("CustomerID", "ProductID")
MINDBOX_HEADERS = (
    "CustomerMindboxId",
    "Quantity",
    "ProductGroupOffline1C",
    "CustomFieldKoefficient",
)
BUYER_HEADERS = (
    "Canonical MindboxID",
    "Телефон",
    "Дисконтная карта",
    "Email",
    "Есть BPR-MF",
    "Проверенных BPR рекомендаций",
    "Гибридных рекомендаций",
    "Историческая вероятность за 30 дней",
    "Качество propensity",
    "Исторический период",
    "Статус рекомендаций",
    "Статус propensity",
)
RAW_HEADERS = (
    "MindboxID",
    "Ранг BPR",
    "Код исходного товара",
    "Название",
    "Коллекция",
    "Raw BPR-score",
    "Допустимость",
)
HYBRID_HEADERS = (
    "MindboxID",
    "Финальный ранг",
    "Код актуального товара",
    "Название",
    "Вид номенклатуры",
    "Коллекция",
    "Остаток",
    "Hybrid final_score",
    "Исходный товар",
    "Название исходного товара",
    "Исходный ранг BPR",
    "Raw BPR-score",
    "Нормализованный BPR-score",
    "Сходство",
    "Вклад ранга",
    "Вклад BPR",
    "Вклад сходства",
    "До бонуса",
    "Бонус коллекции",
    "Делитель",
    "Способ сопоставления",
)
PRODUCT_HEADERS = (
    "Код товара",
    "Название",
    "Вид номенклатуры",
    "Коллекция",
    "Появлений в Top-10",
    "Уникальных покупателей",
    "Средний ранг",
    "Средний final_score",
)
BLUE, PALE = "18324D", "EEF3F8"
ROW_LIMIT = 1_048_576
_BODY_FONT = Font(name="Arial", size=10, color=BLUE)
_HEADER_FONT = Font(name="Arial", size=10, bold=True, color="FFFFFF")
_HEADER_FILL = PatternFill("solid", fgColor=BLUE)
_STRIPE_FILL = PatternFill("solid", fgColor=PALE)
_ALIGNMENTS = {
    (header, text): Alignment(
        vertical="top", wrap_text=True, horizontal="left" if text else "right"
    )
    for header in (False, True)
    for text in (False, True)
}


def _legacy_phone(value):
    # Same normalization as the old BPRMF export, after rejecting extraneous
    # letters/extensions. E.164 validation is reused from current contact search.
    if not isinstance(value, str) or not re.fullmatch(r"\+?[0-9 ()\-]+", value.strip()):
        return None
    digits = re.sub(r"\D+", "", value)
    if len(digits) == 10:
        digits = "7" + digits
    elif len(digits) == 11 and digits.startswith("8"):
        digits = "7" + digits[1:]
    return (
        digits
        if normalize_search_input(ContactType.PHONE, "+" + digits) is not None
        else None
    )


def _code(value):
    return (
        isinstance(value, str)
        and bool(value)
        and value == value.strip()
        and all(c.isprintable() and not c.isspace() and c != "," for c in value)
    )


class _Table:
    def __init__(self, book, name, headers, *, row_limit=ROW_LIMIT, formats=None):
        if not 2 <= row_limit <= ROW_LIMIT:
            raise ExportError("INVALID_ROW_LIMIT")
        self.book, self.name, self.headers, self.limit = book, name, headers, row_limit
        self.formats = formats or {}
        self.parts, self.rows, self.part_rows = [], 0, 0
        self._next()

    def _next(self):
        if self.parts:
            self._finish()
        suffix = "" if not self.parts else f"_{len(self.parts) + 1:02}"
        self.ws = self.book.create_sheet(self.name[: 31 - len(suffix)] + suffix)
        self.parts.append(self.ws.title)
        self.ws.sheet_view.showGridLines = False
        self.ws.sheet_format.defaultRowHeight = 30
        self.ws.row_dimensions[1].height = 60
        self.ws.freeze_panes = "B2"
        for i, title in enumerate(self.headers, 1):
            self.ws.column_dimensions[get_column_letter(i)].width = (
                36
                if "Название" in title or "Качество" in title
                else 27
                if "Статус" in title or "период" in title
                else 23
            )
        if self.name == "06 Методология":
            self.ws.column_dimensions["A"].width = 34
            self.ws.column_dimensions["B"].width = 140
            self.ws.sheet_format.defaultRowHeight = 54
        if self.name == "03 BPR-MF":
            self.ws.column_dimensions["G"].width = 52
        self.ws.append(self._cells(self.headers, header=True))
        self.part_rows = 1

    def _cells(self, values, *, header=False, extra_formats=None):
        for i, value in enumerate(values, 1):
            cell = WriteOnlyCell(self.ws, value=value)
            if isinstance(value, str):
                if len(value) > 32767:
                    raise ExportError("EXPORT_TEXT_TOO_LONG")
                cell.data_type = (
                    "s"  # Literal text, including = + - @; never a formula.
                )
                cell.number_format = "@"
            elif isinstance(value, (float, int, Decimal)) and not isinstance(
                value, bool
            ):
                if not math.isfinite(value):
                    raise ExportError("NONFINITE_EXPORT_NUMBER")
                cell.number_format = (extra_formats or {}).get(
                    i,
                    self.formats.get(
                        i, "#,##0" if isinstance(value, int) else "0.000000"
                    ),
                )
            cell.font = _HEADER_FONT if header else _BODY_FONT
            cell.alignment = _ALIGNMENTS[header, isinstance(value, str)]
            if header:
                cell.fill = _HEADER_FILL
            elif self.rows % 2:
                cell.fill = _STRIPE_FILL
            yield cell

    def append(self, values, *, extra_formats=None):
        if len(values) != len(self.headers):
            raise ExportError("INVALID_EXPORT_ROW")
        if self.part_rows >= self.limit:
            self._next()
        self.ws.append(self._cells(values, extra_formats=extra_formats))
        self.rows += 1
        self.part_rows += 1

    def _finish(self):
        self.ws.auto_filter.ref = (
            f"A1:{get_column_letter(len(self.headers))}{self.part_rows}"
        )

    def finish(self):
        self._finish()


def _quality(p):
    names = {
        "NEAR_ZERO_WITH_RECENT_FAVORITE": "Оценка ненадёжна: недавнее избранное",
        "LOW_INDIVIDUALIZATION": "Недостаточно индивидуальных сигналов",
        "HISTORICAL_FEATURES_UNAVAILABLE": "Неполное покрытие истории",
    }
    return "; ".join(names.get(w, w) for w in p.warnings)


def _validate_bundle(folder, tables, counters, poll):
    with (folder / ONLINE).open(encoding="utf-8-sig", newline="") as stream:
        reader = csv.reader(stream, delimiter=";", strict=True)
        if next(reader) != list(ONLINE_HEADERS):
            raise ExportError("INVALID_ONLINE_CSV")
        count = 0
        for row in reader:
            poll()
            if len(row) != 2 or not row[0] or not row[1]:
                raise ExportError("INVALID_ONLINE_CSV")
            count += 1
        if count != counters["online_rows"]:
            raise ExportError("ROW_COUNT_MISMATCH")
    with (folder / MINDBOX).open(encoding="utf-8-sig", newline="") as stream:
        reader = csv.reader(stream, delimiter=";", strict=True)
        if next(reader) != list(MINDBOX_HEADERS):
            raise ExportError("INVALID_MINDBOX_CSV")
        count = 0
        for row in reader:
            poll()
            if (
                len(row) != 4
                or row[1] != "1"
                or not re.fullmatch(r"-?\d+,\d{2}", row[3])
            ):
                raise ExportError("INVALID_MINDBOX_CSV")
            count += 1
        if count != counters["mindbox_rows"]:
            raise ExportError("ROW_COUNT_MISMATCH")
    book = load_workbook(folder / ANALYTICS, read_only=True, data_only=False)
    try:
        for table in tables:
            for name in table.parts:
                poll()
                if next(book[name].values) != tuple(table.headers):
                    raise ExportError("INVALID_XLSX_HEADER")
    finally:
        book.close()
    # Check every nested XLSX member CRC, without loading cell populations in RAM.
    with zipfile.ZipFile(folder / ANALYTICS) as archive:
        for info in archive.infolist():
            with archive.open(info) as stream:
                while stream.read(1024 * 1024):
                    poll()


class _PolledFile:
    """Cooperative cancellation/resource checks during openpyxl ZIP compression."""

    def __init__(self, stream, poll):
        self.stream, self.poll, self.failed = stream, poll, False

    def write(self, data):
        if not self.failed:
            try:
                self.poll()
            except BaseException:
                self.failed = True  # Allow zipfile's exception cleanup to close.
                raise
        return self.stream.write(data)

    def __getattr__(self, name):
        return getattr(self.stream, name)


def export_recommendations(
    destination,
    records,
    metadata,
    *,
    cancelled=lambda: False,
    guard=lambda: None,
    progress=lambda **kw: None,
    overwrite=False,
    row_limit=ROW_LIMIT,
    disk_floor=20 * 1024**3,
    protect=None,
    monitor=lambda: None,
):
    """Private staging, full verification, atomic publication; no partial target.

    Call only inside the isolated CLI export process: temporary XML redirection
    intentionally uses process-local tempfile.tempdir and is restored on exit.
    """
    started = perf_counter()
    destination = Path(destination).absolute()
    if destination.name != "Recommendations.zip" or not destination.parent.is_dir():
        raise ExportError("INVALID_DESTINATION")
    if destination.exists() and not overwrite:
        raise ExportError("DESTINATION_EXISTS")
    if destination.is_symlink():
        raise ExportError("UNSAFE_DESTINATION")
    if protect is None:
        from Application.model.propensity_ml.resources import protect_directory

        protect = protect_directory
    # Estimate uncompressed sheet XML + final XLSX/ZIP staging, keep independent reserve.
    required = max(64 * 1024**2, metadata.total_candidates * 35_000)
    if shutil.disk_usage(destination.parent).free < disk_floor + required:
        raise ExportError("INSUFFICIENT_DISK")
    folder = Path(tempfile.mkdtemp(prefix=".recommendations-", dir=destination.parent))
    old_temp = tempfile.tempdir
    book = None
    tables = []
    stage_db = None
    iterator = None
    counters = Counter()
    products = {}
    collections = Counter()
    kinds = Counter()
    score_bins = Counter()
    last_poll = 0.0

    def poll():
        nonlocal last_poll
        if cancelled():
            raise ExportCancelled
        now = perf_counter()
        if now - last_poll >= 1:
            monitor()
            if shutil.disk_usage(folder).free < disk_floor:
                raise ExportError("INSUFFICIENT_DISK")
            last_poll = now

    try:
        protect(folder)  # No personal data is written before permissions succeed.
        if os.name != "nt":
            folder.chmod(0o700)
        tempfile.tempdir = str(folder)
        iterator = iter(records)
        book = Workbook(write_only=True)
        summary = book.create_sheet("01 Сводка")
        buyers = _Table(
            book,
            "02 Покупатели",
            BUYER_HEADERS,
            row_limit=row_limit,
            formats={8: "0.00%"},
        )
        raw_sheet = _Table(
            book, "03 BPR-MF", RAW_HEADERS, row_limit=row_limit, formats={6: "0.000000"}
        )
        hybrid_sheet = _Table(
            book,
            "04 Гибридные рекомендации",
            HYBRID_HEADERS,
            row_limit=row_limit,
            formats={7: "#,##0.00", 8: "0.000000"},
        )
        product_sheet = _Table(
            book,
            "05 Аналитика товаров",
            PRODUCT_HEADERS,
            row_limit=row_limit,
            formats={7: "0.00", 8: "0.000000"},
        )
        method = _Table(
            book, "06 Методология", ("Параметр", "Значение"), row_limit=row_limit
        )
        tables = [buyers, raw_sheet, hybrid_sheet, product_sheet, method]
        stage_db = sqlite3.connect(folder / "online-stage.sqlite")
        stage_db.execute("PRAGMA cache_size=-4096")
        stage_db.execute("PRAGMA temp_store=FILE")
        stage_db.execute("CREATE TABLE phones(phone TEXT, products TEXT)")
        stage_db.execute("CREATE INDEX phone_sort ON phones(phone,products)")
        last_id = None
        with (folder / MINDBOX).open("w", encoding="utf-8-sig", newline="") as stream:
            writer = csv.writer(stream, delimiter=";")
            writer.writerow(MINDBOX_HEADERS)
            for record in iterator:
                poll()
                counters["checked"] += 1
                if record is None:
                    counters["skipped_profiles"] += 1
                    if counters["checked"] % 100 == 0:
                        progress(
                            phase="customers",
                            processed=counters["checked"],
                            total=metadata.total_candidates,
                        )
                    continue
                customer = record.canonical_id
                if (
                    not isinstance(customer, str)
                    or not customer
                    or customer != customer.strip()
                    or any(not c.isprintable() or c.isspace() for c in customer)
                    or (last_id is not None and customer <= last_id)
                ):
                    raise ExportError("DUPLICATE_OR_UNORDERED_CUSTOMER")
                last_id = customer
                counters["customers"] += 1
                result, p = record.hybrid, record.propensity
                available = p.status.value in ("READY", "WARNING")
                probability = p.probability if available else None
                period = (
                    f"{p.prediction_date[:10]} — {p.target_until[:10]}"
                    if available
                    else None
                )
                quality = _quality(p) if available else None
                counters["with_propensity"] += int(available)
                counters["propensity_" + p.status.value] += 1
                raw = (
                    tuple(
                        rec
                        for rec in result.displayed_raw
                        if _code(rec.model_item_code)
                    )
                    if result and result.raw_status in ("READY", "EMPTY")
                    else ()
                )
                items = []
                seen = set()
                for item in (
                    result.items[:10] if result and result.status == "READY" else ()
                ):
                    if (
                        not _code(item.target_item_code)
                        or item.target_item_code in seen
                    ):
                        continue
                    if not math.isfinite(item.final_score):
                        raise ExportError("NONFINITE_SCORE")
                    seen.add(item.target_item_code)
                    items.append(item)
                counters["with_hybrid"] += int(bool(items))
                counters["with_bpr"] += int(
                    result is not None
                    and result.raw_bpr is not None
                    and result.raw_bpr.status.value == "OK"
                )
                counters["status_" + record.data_status] += 1
                c = record.contact
                buyers.append(
                    (
                        customer,
                        c.mobile_phone,
                        c.discount_card,
                        c.email,
                        bool(
                            result
                            and result.raw_bpr
                            and result.raw_bpr.status.value == "OK"
                        ),
                        len(raw),
                        len(items),
                        probability,
                        quality,
                        period,
                        record.data_status,
                        p.status.value,
                    ),
                    extra_formats={8: "0.000E+00"}
                    if available and "NEAR_ZERO_WITH_RECENT_FAVORITE" in p.warnings
                    else None,
                )
                names = (
                    {item.code: item for item in result.historical_items}
                    if result
                    else {}
                )
                for rec in raw:
                    if not _code(rec.model_item_code):
                        continue
                    info = names.get(rec.model_item_code)
                    raw_sheet.append(
                        (
                            customer,
                            rec.rank,
                            rec.model_item_code,
                            info.name if info else None,
                            info.collection if info else None,
                            rec.raw_score,
                            "Проверен; разрешённый вид; не куплен и не UNCERTAIN",
                        )
                    )
                if (
                    items
                    and record.phone_confirmed is True
                    and record.phone_invalid is not True
                ):
                    phone = _legacy_phone(c.mobile_phone)
                    if phone:
                        stage_db.execute(
                            "INSERT INTO phones VALUES(?,?)",
                            (phone, ",".join(i.target_item_code for i in items)),
                        )
                for item in items:
                    if len(item.sources) != 1:
                        raise ExportError("INCOMPATIBLE_SCORE_PROVENANCE")
                    source = item.sources[0]
                    calc = source.calculation

                    def value(name):
                        return getattr(calc, name, None)

                    writer.writerow(
                        (
                            customer,
                            1,
                            item.target_item_code,
                            f"{item.final_score:.2f}".replace(".", ","),
                        )
                    )
                    counters["mindbox_rows"] += 1
                    hybrid_sheet.append(
                        (
                            customer,
                            item.final_rank,
                            item.target_item_code,
                            item.target_name,
                            item.target_kind,
                            item.target_collection,
                            item.target_stock,
                            item.final_score,
                            source.source_model_item_code,
                            source.source_name,
                            source.source_raw_rank,
                            source.source_raw_score,
                            value("normalized_bpr_score"),
                            source.similarity_score,
                            value("rank_contribution"),
                            value("bpr_contribution"),
                            value("similarity_contribution"),
                            value("before_bonus"),
                            value("bonus"),
                            value("divisor"),
                            source.match_method,
                        )
                    )
                    aggregate = products.setdefault(
                        item.target_item_code,
                        [
                            item.target_name,
                            item.target_kind,
                            item.target_collection,
                            0,
                            0.0,
                            0.0,
                        ],
                    )
                    aggregate[3] += 1
                    aggregate[4] += item.final_rank
                    aggregate[5] += item.final_score
                    collections[item.target_collection] += 1
                    kinds[item.target_kind] += 1
                    score_bins[min(4, max(0, int(item.final_score * 5)))] += 1
                if counters["checked"] % 100 == 0:
                    stage_db.commit()
                    progress(
                        phase="customers",
                        processed=counters["checked"],
                        total=metadata.total_candidates,
                    )
        stage_db.commit()
        with (folder / ONLINE).open("w", encoding="utf-8-sig", newline="") as stream:
            writer = csv.writer(stream, delimiter=";")
            writer.writerow(ONLINE_HEADERS)
            # A shared phone cannot identify one customer: exclude every collision.
            for phone, count, items in stage_db.execute(
                "SELECT phone,COUNT(*),MIN(products) FROM phones GROUP BY phone ORDER BY phone"
            ):
                poll()
                if count != 1:
                    counters["phone_collisions"] += count
                    continue
                writer.writerow((phone, items))
                counters["online_rows"] += 1
        stage_db.close()
        stage_db = None
        for code, (name, kind, collection, count, rank, score) in sorted(
            products.items()
        ):
            product_sheet.append(
                (
                    code,
                    name,
                    kind,
                    collection,
                    count,
                    count,
                    rank / count,
                    score / count,
                )
            )
        for label, val in metadata.provenance:
            method.append((label, val))
        exported_at = (
            datetime.now(timezone.utc)
            .astimezone(timezone(timedelta(hours=3)))
            .isoformat()
        )
        for row in (
            ("Export date (Europe/Moscow)", exported_at),
            (
                "Propensity",
                "Исторический RESEARCH_ONLY; срез 01.09.2026, горизонт [01.09.2026, 01.10.2026). Не актуальный прогноз на дату экспорта.",
            ),
            (
                "Propensity quality",
                "Индивидуальные оценки могут быть неточны; near-zero с избранным ненадёжен. Сохранённое число — доля в научной записи при near-zero; отсутствие — пустая ячейка.",
            ),
            (
                "BPR-score",
                "Исходная модельная оценка; не вероятность. Только проверенные displayed_raw, как в карточке.",
            ),
            (
                "Hybrid final_score",
                "Готовое HybridItem.final_score; формула при экспорте не пересчитывается. Это рейтинг, не вероятность.",
            ),
            (
                "Sources",
                "По одному результату на подтверждённого canonical клиента; единые товары/порядок/оценки для всех трёх файлов.",
            ),
            (
                "Online CSV",
                "CustomerID;ProductID, UTF-8 BOM; подтверждённый неinvalid телефон, legacy digits normalization; общие телефоны аудитории исключены.",
            ),
            (
                "Mindbox CSV",
                "CustomerMindboxId;Quantity;ProductGroupOffline1C;CustomFieldKoefficient; Quantity=1, final_score с двумя знаками и запятой.",
            ),
            (
                "Exclusions",
                "Покупки/UNCERTAIN, виды/коллекции/остатки — существующие правила карточки. Пустые, повреждённые и повторные коды не экспортируются.",
            ),
            (
                "Product analytics",
                "По фактически экспортированным hybrid строкам. Одна уникальная пара покупатель/товар; появления равны числу уникальных покупателей.",
            ),
            (
                "Source changes",
                "Проверки на границах пакетов и перед публикацией. Любое изменение отменяет весь неопубликованный архив.",
            ),
        ):
            method.append(row)
        summary.sheet_view.showGridLines = False
        summary.column_dimensions["A"].width = 60
        summary.column_dimensions["B"].width = 26
        title = WriteOnlyCell(summary, value="KANZLER AI — рекомендации")
        title.font = Font(name="Arial", size=16, bold=True, color="FFFFFF")
        title.fill = _HEADER_FILL
        summary.append([title])
        summary.append([])

        def summary_row(label, val):
            cells = []
            for item in (label, val):
                cell = WriteOnlyCell(summary, value=item)
                if isinstance(item, str):
                    cell.data_type = "s"
                cell.font = _BODY_FONT
                cells.append(cell)
            summary.append(cells)

        for label, val in (
            ("Проверено покупателей", counters["checked"]),
            ("Подтверждённых покупателей", counters["customers"]),
            ("С гибридными рекомендациями", counters["with_hybrid"]),
            ("С BPR-MF", counters["with_bpr"]),
            ("С историческим propensity", counters["with_propensity"]),
            (
                "Без гибридных рекомендаций",
                counters["customers"] - counters["with_hybrid"],
            ),
            ("Строк Online Store CSV", counters["online_rows"]),
            ("Строк Mindbox CSV", counters["mindbox_rows"]),
            ("Уникальных актуальных товаров", len(products)),
            (
                "Среднее рекомендаций на подтверждённого клиента",
                counters["mindbox_rows"] / counters["customers"]
                if counters["customers"]
                else 0,
            ),
            ("Исключены: профиль не подтверждён", counters["skipped_profiles"]),
            ("Исключены из Online: общий телефон", counters["phone_collisions"]),
        ):
            summary_row(label, val)
        summary.append([])
        summary_row("Диапазон final_score (описательный)", "Число рекомендаций")
        score_start = 17
        for i in range(5):
            summary_row(f"{i / 5:.1f}–{(i + 1) / 5:.1f}", score_bins[i])
        chart = BarChart()
        chart.title = "Распределение итоговых оценок"
        chart.y_axis.title = "Рекомендации"
        chart.add_data(
            Reference(summary, min_col=2, min_row=score_start, max_row=score_start + 4)
        )
        chart.set_categories(
            Reference(summary, min_col=1, min_row=score_start, max_row=score_start + 4)
        )
        chart.style = 13
        chart.legend = None
        chart.series[0].graphicalProperties.solidFill = BLUE
        chart.series[0].graphicalProperties.line.solidFill = BLUE
        summary.add_chart(chart, "D3")
        summary.append([])
        summary_row("Коллекция", "Рекомендаций")
        for label, count in sorted(collections.items()):
            summary_row(label, count)
        summary.append([])
        summary_row("Вид номенклатуры", "Рекомендаций")
        for label, count in sorted(kinds.items()):
            summary_row(label, count)
        summary.append([])
        summary_row("Дата экспорта (Europe/Moscow)", exported_at)
        summary_row("Аудитория", metadata.audience)
        summary_row("Всего клиентов BPR-MF в модели", metadata.model_users)
        summary_row("Строк совместимого propensity-артефакта", metadata.propensity_rows)
        for table in tables:
            table.finish()
        progress(
            phase="xlsx", processed=counters["checked"], total=metadata.total_candidates
        )
        poll()
        guard()
        with (folder / ANALYTICS).open("w+b") as stream:
            with zipfile.ZipFile(
                _PolledFile(stream, poll), "w", zipfile.ZIP_DEFLATED, allowZip64=True
            ) as archive:
                ExcelWriter(book, archive).save()
        counters["raw_rows"] = raw_sheet.rows
        counters["hybrid_rows"] = hybrid_sheet.rows
        progress(
            phase="verify",
            processed=counters["checked"],
            total=metadata.total_candidates,
        )
        _validate_bundle(folder, tables, counters, poll)
        temporary = folder / "complete.zip"
        with zipfile.ZipFile(
            temporary, "w", compression=zipfile.ZIP_DEFLATED, compresslevel=6
        ) as archive:
            for name in MEMBERS:
                info = zipfile.ZipInfo(name)
                info.compress_type = (
                    zipfile.ZIP_STORED if name == ANALYTICS else zipfile.ZIP_DEFLATED
                )
                with (
                    (folder / name).open("rb") as source,
                    archive.open(info, "w", force_zip64=True) as target,
                ):
                    while chunk := source.read(1024 * 1024):
                        poll()
                        target.write(chunk)
        with zipfile.ZipFile(temporary) as archive:
            if tuple(archive.namelist()) != MEMBERS:
                raise ExportError("INVALID_ZIP_MEMBERS")
            for info in archive.infolist():
                with archive.open(info) as stream:
                    while stream.read(1024 * 1024):
                        poll()
        poll()
        guard()
        with temporary.open("r+b") as stream:
            os.fsync(stream.fileno())
        size = temporary.stat().st_size
        if overwrite:
            os.replace(temporary, destination)
        elif os.name == "nt":
            os.rename(
                temporary, destination
            )  # Atomic, refuses an existing destination on Windows.
        else:
            os.link(temporary, destination)
            temporary.unlink()
        return ExportSummary(
            counters["checked"],
            counters["customers"],
            counters["online_rows"],
            counters["mindbox_rows"],
            raw_sheet.rows,
            hybrid_sheet.rows,
            size,
            perf_counter() - started,
            counters["skipped_profiles"],
            counters["phone_collisions"],
            {
                k: v
                for k, v in counters.items()
                if k.startswith(("status_", "propensity_"))
            },
        )
    finally:
        close = getattr(iterator, "close", None)
        if callable(close):
            close()
        if stage_db is not None:
            stage_db.close()
        if book is not None:
            for sheet in book.worksheets:
                if not sheet.closed:
                    try:
                        sheet.close()
                    except Exception:
                        pass
                writer = getattr(sheet, "_writer", None)
                if writer is not None and getattr(writer, "out", None):
                    Path(writer.out).unlink(missing_ok=True)
            book.close()
        tempfile.tempdir = old_temp
        shutil.rmtree(folder)
