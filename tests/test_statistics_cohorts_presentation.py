"""Saved cohort snapshot parity across actual GUI, XLSX and offline HTML."""

from copy import deepcopy
from datetime import datetime
from pathlib import Path
import json
import sqlite3

import pytest
from openpyxl import load_workbook
from PyQt6.QtWidgets import QApplication, QLabel, QTableWidget, QWidget, QVBoxLayout

from Application import dataset_statistics as statistics, statistics_cache as cache, statistics_slices as slices
from Application import statistics_export as xlsx, statistics_html_export as html
from Application.analysis_filter import AnalysisFilter
from Application.statistics_cohorts_core import CohortAggregates
from Application.statistics_cohorts_presentation import cohort_sections, COHORT_TITLE, COHORT_LEGACY_TEXT, MATRIX_HEADERS
from Application.tabs import dataset_statistics_tab as ui
from tests.test_statistics_cohorts_integration import dataset, calculate, SELECTION, save
from tests.tabs.test_dataset_statistics_tab import window, app


__all__ = ['dataset','window','app']
pytestmark = pytest.mark.usefixtures('window_settings')


def _snapshot(dataset, scenario):
    if scenario=='no_orders':
        dataset[2]['manual_interactions'] = None
        dataset[2]['orders'] = {}
        save(dataset[0],dataset[2])
        return calculate(dataset,AnalysisFilter())
    if scenario=='partial':
        manual = dataset[2]['manual_interactions']
        for target in (manual,manual['orders'],manual['actions']):
            target.update(since='2026-01-01T00:00:00Z',until='2026-01-10T00:00:00Z')
        save(dataset[0],dataset[2])
    result = calculate(dataset,AnalysisFilter('2025-01-01','2025-08-31') if scenario=='empty' else SELECTION)
    if scenario=='legacy':
        result.pop('cohort_statistics')
    return result


@pytest.mark.parametrize('scenario',['full','partial','empty','no_orders','legacy'])
def test_gui_xlsx_html_show_same_saved_tables_without_calculation_or_sources(dataset,window,tmp_path,monkeypatch,scenario):
    result = _snapshot(dataset,scenario)
    path = tmp_path/'cohort-result.json'
    cache.save_result(path,result)
    result = cache.load_result(path)
    before = deepcopy(result)
    saved = path.read_bytes()
    _,tab = window
    def forbidden(*args,**kwargs):
        raise AssertionError('Cohort presentation attempted source access or calculation')
    with monkeypatch.context() as patch:
        patch.setattr(statistics,'iter_export',forbidden)
        patch.setattr(statistics,'calculate_statistics',forbidden)
        patch.setattr(slices,'calculate_cached_statistics',forbidden)
        patch.setattr(CohortAggregates,'add',forbidden)
        patch.setattr(CohortAggregates,'result',forbidden)
        patch.setattr(sqlite3,'connect',forbidden)
        original_open = Path.open
        def guarded_open(p,*args,**kwargs):
            if p.resolve().is_relative_to(ui.PROJECT_ROOT/'input_data'):
                forbidden()
            return original_open(p,*args,**kwargs)
        patch.setattr(Path,'open',guarded_open)
        tab.render(result)
        tab.sections.setCurrentIndex(3)
        QApplication.processEvents()
        target = tmp_path/'cohort.xlsx'
        xlsx.export_statistics(result,target)
        rendered = html.build_statistics_html(result)
    assert result == before and path.read_bytes() == saved
    sections = cohort_sections(result)
    page = tab.sections.widget(3)
    labels = [label.text() for label in page.findChildren(QLabel)]
    assert all(title.upper() in labels for title,_,_,_ in sections)
    expected = [(headers,rows) for _,_,blocks,_ in sections for headers,rows in blocks]
    tables = page.findChildren(QTableWidget)[-len(expected):]
    assert len(tables) == len(expected)
    for table,(headers,rows) in zip(tables,expected):
        assert [table.horizontalHeaderItem(c).text() for c in range(table.columnCount())] == list(headers)
        assert [[table.item(r,c).text() for c in range(table.columnCount())] for r in range(table.rowCount())] == [
            [ui._number(ui._buyer_display(v)) for v in row] for row in rows]
        assert all(table.rowHeight(r) >= table.sizeHintForRow(r) for r in range(table.rowCount()))
        if headers == list(MATRIX_HEADERS) and rows:
            # The legacy test window can expand enough to fit every column.
            # Exercise the matrix itself at a narrow viewport explicitly.
            table.setMaximumWidth(900)
            QApplication.processEvents()
            assert table.horizontalScrollBar().maximum() > 0
            assert table.font().pointSize() >= 9
    book = load_workbook(target)
    try:
        assert book.sheetnames == list(xlsx.SHEET_NAMES)
        rows = list(book['Когортный анализ'].values)
        html_blocks = html.presentation(result)[4].blocks
        all_html_rows = [r for block in html_blocks if block['kind']=='table' for r in block['rows']]
        for title,cards,blocks,notes in cohort_sections(result,include_extended=True):
            assert any(r[0]==title for r in rows)
            assert html.display_text(title) in rendered
            for label,value in cards:
                matching = [r for r in rows if r[0]==label]
                assert matching
                actual = matching[0][1]
                assert (actual.date() if isinstance(actual,datetime) else actual)==value
            for headers,values in blocks:
                assert any(tuple(r[:len(headers)])==tuple(headers) for r in rows)
                for value_row in values:
                    assert any([c['text'] for c in r]==[html._cell(v)['text'] for v in value_row] for r in all_html_rows)
                    assert any(all(actual==pytest.approx(v) if isinstance(v,float) else actual==v
                                   for actual,v in zip(candidate,value_row)) for candidate in rows)
            assert all(html._cell(text)['text'] in rendered for _,text in notes)
        if scenario=='legacy':
            assert COHORT_LEGACY_TEXT in rendered
        if scenario=='full':
            # M1 is a numeric percentage; future M12 remains text, not zero.
            sheet = book['Когортный анализ']
            header_row = next(i for i,r in enumerate(rows,1) if tuple(r[:14])==MATRIX_HEADERS)
            assert sheet.cell(header_row+1,3).data_type=='n'
            assert '%' in sheet.cell(header_row+1,3).number_format
            assert sheet.cell(header_row+1,14).value=='Нет данных'
            assert 'table-scroll' in rendered and 'M12, %' in rendered
        content = json.dumps(result.get('cohort_statistics'))
        assert all(key not in content for key in ('customer_id','order_id','identity_space','secret-'))
        assert 'secret-historical' not in rendered
    finally:
        book.close()


def test_legacy_cache_opens_on_startup_without_automatic_recalculation(dataset,window,tmp_path,monkeypatch):
    result = calculate(dataset)
    result.pop('cohort_statistics')
    path = tmp_path/'old-cohortless.json'
    cache.save_result(path,result)
    before = path.read_bytes()
    monkeypatch.setattr(ui,'load_filter',lambda:SELECTION)
    monkeypatch.setattr(ui,'CACHE_PATH',path)
    monkeypatch.setattr(ui.DatasetStatisticsTab,'start',lambda self:pytest.fail('Automatic recalculation'))
    tab = ui.DatasetStatisticsTab(window[0])
    try:
        assert tab.displayed_result is not None and 'cohort_statistics' not in tab.displayed_result
        assert path.read_bytes()==before
        assert any(label.text()==COHORT_TITLE.upper() for label in tab.findChildren(QLabel))
        tables = tab.sections.widget(3).findChildren(QTableWidget)
        assert tables[-1].item(0,1).text()==COHORT_LEGACY_TEXT
    finally:
        tab._shutdown()
        tab.deleteLater()


def test_extended_matrices_remain_in_exports_without_duplicating_gui_matrix(dataset):
    result = calculate(dataset)
    gui = cohort_sections(result)
    exports = cohort_sections(result,include_extended=True)
    assert exports[:len(gui)] == gui
    assert sum('матрица удержания' in title for title,_,_,_ in gui)==1
    assert sum('матрица удержания' in title for title,_,_,_ in exports)==2



def test_gui_cohort_rounding_matches_html_without_rounding_stored_values(app,monkeypatch):
    number = 2.675  # Binary float formatting with .2f alone would produce 2.67.
    monkeypatch.setattr(ui,'buyer_sections',lambda result:[])
    monkeypatch.setattr(ui,'rfm_sections',lambda result:[])
    monkeypatch.setattr(ui,'cohort_sections',lambda result:[(COHORT_TITLE,[],[(['M1, %'],[(number,)])],[])])
    widget = QWidget()
    try:
        ui._buyer_page(QVBoxLayout(widget),{})
        table = widget.findChild(QTableWidget)
        assert table.item(0,0).text() == html._cell(number)['text'].replace(',','.') == '2.68'
        assert number == 2.675  # Presentation did not change the saved value.
    finally:
        widget.close()
        widget.deleteLater()
