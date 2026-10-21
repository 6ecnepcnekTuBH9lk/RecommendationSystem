"""Synthetic one-stream ZIP contracts, transactional publication and Excel limits."""

from dataclasses import replace
from decimal import Decimal
import csv
import io
from types import SimpleNamespace
import zipfile

from openpyxl import load_workbook
import pytest

from Application import recommendation_export as writer
from Application.customer_profiles import CustomerContact
from Application.customer_propensity_service import (
    CustomerPropensity,
    PropensityStatus as P,
)
from Application.hybrid_recommendations import (
    HybridResult,
    HybridItem,
    MatchSource,
    HistoricalItem,
)
from Application.model.raw_bpr_inference import (
    RawBprRecommendation,
    RawBprResult,
    RawBprStatus,
)
from Application.recommendation_export_records import (
    ExportCustomer,
    ExportMetadata,
    ExportError,
    ExportCancelled,
)


def customer(
    number=1,
    *,
    count=3,
    phone="+7 (916) 123-45-67",
    confirmed=True,
    probability=0.123456,
):
    raw = tuple(
        RawBprRecommendation(i + 1, f"S{i + 1}", 9.123456789 - i) for i in range(count)
    )
    items = tuple(
        HybridItem(
            f"000{i + 1:03}",
            f"Товар {i + 1}",
            "Брюки",
            "Весна-Лето 2026",
            Decimal("12.50"),
            "",
            i + 1,
            0.98765 - i * 0.1,
            (
                MatchSource(
                    f"S{i + 1}",
                    f"Исторический товар {i + 1}",
                    raw[i].raw_score,
                    i + 1,
                    0.75,
                    "weighted_attributes_v1",
                    0.5,
                ),
            ),
        )
        for i in range(count)
    )
    hybrid = HybridResult(
        "READY" if items else "EMPTY",
        "",
        items,
        RawBprResult("g", f"M{number}", RawBprStatus.OK, raw),
        historical_items=tuple(
            HistoricalItem(r.model_item_code, f"Источник {i}", "2025")
            for i, r in enumerate(raw)
        ),
        displayed_raw=raw,
        raw_status="READY" if raw else "EMPTY",
    )
    propensity = (
        CustomerPropensity(
            P.READY,
            f"C{number}",
            probability,
            "2026-09-01T00:00:00Z",
            "2026-10-01T00:00:00Z",
            "a" * 64,
        )
        if probability is not None
        else CustomerPropensity(P.NOT_FOUND, f"C{number}")
    )
    return ExportCustomer(
        f"C{number}",
        CustomerContact("mail@example.invalid", phone, "0000123"),
        confirmed,
        False,
        hybrid,
        propensity,
    )


@pytest.fixture
def export(tmp_path):
    def run(records, **kwargs):
        records = list(records)
        path = tmp_path / "Recommendations.zip"
        result = writer.export_recommendations(
            path,
            records,
            ExportMetadata("synthetic", len(records), 2, 2),
            protect=lambda p: None,
            disk_floor=0,
            **kwargs,
        )
        return path, result

    return run


def read_csv(archive, name):
    payload = archive.read(name)
    assert payload.startswith(b"\xef\xbb\xbf")
    return list(csv.reader(io.StringIO(payload.decode("utf-8-sig")), delimiter=";"))


def test_exact_files_csv_schemas_scores_and_consistent_order(export, tmp_path):
    path, result = export([customer()])
    with zipfile.ZipFile(path) as archive:
        assert tuple(archive.namelist()) == writer.MEMBERS
        online = read_csv(archive, writer.ONLINE)
        mindbox = read_csv(archive, writer.MINDBOX)
        assert online == [
            list(writer.ONLINE_HEADERS),
            ["79161234567", "000001,000002,000003"],
        ]
        assert mindbox[0] == list(writer.MINDBOX_HEADERS)
        assert all(len(r) == 4 for r in mindbox)
        assert [r[2] for r in mindbox[1:]] == online[1][1].split(",")
        assert [r[3] for r in mindbox[1:]] == ["0,99", "0,89", "0,79"]
        assert all(r[:2] == ["C1", "1"] for r in mindbox[1:])
        xlsx = tmp_path / "inspect.xlsx"
        xlsx.write_bytes(archive.read(writer.ANALYTICS))
    book = load_workbook(xlsx, read_only=False)
    try:
        assert len(book.sheetnames) == 6
        rows = list(book["04 Гибридные рекомендации"].values)
        assert rows[0] == writer.HYBRID_HEADERS
        assert [r[2] for r in rows[1:]] == online[1][1].split(",")
        assert [r[7] for r in rows[1:]] == [
            i.final_score for i in customer().hybrid.items
        ]
        assert book["02 Покупатели"]["A2"].data_type == "s"
        assert book["02 Покупатели"]["C2"].value == "0000123"
        assert book["02 Покупатели"]["H2"].value == 0.123456
        assert book["02 Покупатели"]["H2"].number_format == "0.00%"
        assert book["02 Покупатели"].freeze_panes == "B2"
        assert book["04 Гибридные рекомендации"].auto_filter.ref == "A1:U4"
        assert book["03 BPR-MF"]["F2"].value == 9.123456789
        assert len(book["01 Сводка"]._charts) == 1
    finally:
        book.close()
    assert result.online_rows == 1 and result.mindbox_rows == result.hybrid_rows == 3
    assert not list(path.parent.glob(".recommendations-*"))


@pytest.mark.parametrize(
    "phone,expected",
    [
        ("+7 (916) 123-45-67", "79161234567"),
        ("8 916 123-45-67", "79161234567"),
        ("9161234567", "79161234567"),
        ("+44 20 7946 0958", "442079460958"),
        ("phone 79161234567", None),
        ("123", None),
    ],
)
def test_existing_legacy_phone_format(phone, expected):
    assert writer._legacy_phone(phone) == expected


@pytest.mark.parametrize(
    "phone,confirmed", [(None, True), ("+79161234567", False), ("+79161234567", None)]
)
def test_no_phone_or_unconfirmed_does_not_block_mindbox(export, phone, confirmed):
    path, result = export([customer(phone=phone, confirmed=confirmed)])
    with zipfile.ZipFile(path) as archive:
        assert len(read_csv(archive, writer.ONLINE)) == 1
        assert len(read_csv(archive, writer.MINDBOX)) == 4
    assert result.online_rows == 0 and result.mindbox_rows == 3


def test_phone_collisions_exclude_every_online_row(export):
    path, result = export([customer(1), customer(2)])
    with zipfile.ZipFile(path) as archive:
        assert len(read_csv(archive, writer.ONLINE)) == 1
    assert result.excluded_phone_collisions == 2 and result.mindbox_rows == 6


def test_missing_propensity_is_blank_not_zero(export, tmp_path):
    path, _ = export([customer(probability=None)])
    with zipfile.ZipFile(path) as archive:
        payload = tmp_path / "analytics.xlsx"
        payload.write_bytes(archive.read(writer.ANALYTICS))
    book = load_workbook(payload, read_only=True)
    try:
        row = list(book["02 Покупатели"].values)[1]
        assert row[7] is None and row[9] is None and row[11] == "NOT_FOUND"
    finally:
        book.close()


def test_split_sheets_without_lost_rows_and_literal_formula_text(export, tmp_path):
    c = customer(count=7)
    items = tuple(replace(i, target_name='=HYPERLINK("bad")') for i in c.hybrid.items)
    c = replace(c, hybrid=replace(c.hybrid, items=items))
    path, result = export([c], row_limit=4)
    with zipfile.ZipFile(path) as archive:
        payload = tmp_path / "analytics.xlsx"
        payload.write_bytes(archive.read(writer.ANALYTICS))
    book = load_workbook(payload, read_only=False)
    try:
        sheets = [s for s in book if s.title.startswith("04 ")]
        assert len(sheets) == 3 and sum(s.max_row - 1 for s in sheets) == 7
        assert all(
            s["D2"].data_type == "s" and s["D2"].value.startswith("=HYPERLINK")
            for s in sheets
        )
        assert all(s.max_row <= 4 for s in sheets)
    finally:
        book.close()
    assert result.hybrid_rows == 7


def test_duplicates_and_bad_codes_not_exported(export):
    c = customer()
    c = replace(
        c,
        hybrid=replace(
            c.hybrid,
            items=(
                c.hybrid.items[0],
                c.hybrid.items[0],
                replace(c.hybrid.items[1], target_item_code=""),
            ),
        ),
    )
    path, result = export([c])
    assert result.hybrid_rows == 1
    with zipfile.ZipFile(path) as archive:
        assert read_csv(archive, writer.ONLINE)[1][1] == "000001"


def test_existing_destination_never_overwritten_without_optin(export, tmp_path):
    path = tmp_path / "Recommendations.zip"
    path.write_bytes(b"keep existing")
    with pytest.raises(ExportError, match="DESTINATION_EXISTS"):
        export([customer()])
    assert path.read_bytes() == b"keep existing"


@pytest.mark.parametrize(
    "failure", [ExportCancelled(), ExportError("SOURCES_CHANGED"), OSError("disk full")]
)
def test_failure_leaves_no_archive_or_staging(export, tmp_path, failure):
    def guard():
        raise failure

    with pytest.raises(type(failure)):
        export([customer()], guard=guard)
    assert not (tmp_path / "Recommendations.zip").exists()
    assert not list(tmp_path.glob(".recommendations-*"))


def test_cancel_during_stream_closes_generator_and_cleans(tmp_path):
    state = {"cancelled": False, "closed": False}

    def records():
        try:
            yield customer(1)
            state["cancelled"] = True
            yield customer(2)
        finally:
            state["closed"] = True

    with pytest.raises(ExportCancelled):
        writer.export_recommendations(
            tmp_path / "Recommendations.zip",
            records(),
            ExportMetadata("bpr", 2, 2, 0),
            cancelled=lambda: state["cancelled"],
            protect=lambda p: None,
            disk_floor=0,
        )
    assert state["closed"]
    assert not (tmp_path / "Recommendations.zip").exists()
    assert not list(tmp_path.glob(".recommendations-*"))


def test_cancel_during_final_verification_preserves_existing_archive(export, tmp_path):
    path = tmp_path / "Recommendations.zip"
    path.write_bytes(b"old")
    state = {"cancelled": False}

    def progress(**payload):
        if payload["phase"] == "verify":
            state["cancelled"] = True

    with pytest.raises(ExportCancelled):
        export(
            [customer()],
            overwrite=True,
            progress=progress,
            cancelled=lambda: state["cancelled"],
        )
    assert path.read_bytes() == b"old" and not list(tmp_path.glob(".recommendations-*"))


def test_near_zero_is_numeric_scientific_with_quality_warning(export):
    c = customer(probability=1e-30)
    c = replace(
        c,
        propensity=replace(
            c.propensity,
            status=P.WARNING,
            warnings=("NEAR_ZERO_WITH_RECENT_FAVORITE", "LOW_INDIVIDUALIZATION"),
        ),
    )
    path, _ = export([c])
    with zipfile.ZipFile(path) as archive:
        book = load_workbook(io.BytesIO(archive.read(writer.ANALYTICS)))
    try:
        assert book["02 Покупатели"]["H2"].value == 1e-30
        assert book["02 Покупатели"]["H2"].number_format == "0.000E+00"
        assert "ненадёжна" in book["02 Покупатели"]["I2"].value
        assert "Недостаточно" in book["02 Покупатели"]["I2"].value
    finally:
        book.close()


def test_disk_reserve_fails_before_personal_data(tmp_path, monkeypatch):
    monkeypatch.setattr(writer.shutil, "disk_usage", lambda p: SimpleNamespace(free=1))
    with pytest.raises(ExportError, match="INSUFFICIENT_DISK"):
        writer.export_recommendations(
            tmp_path / "Recommendations.zip",
            [customer()],
            ExportMetadata("bpr", 1, 1, 1),
            protect=lambda p: None,
        )
    assert not list(tmp_path.glob(".recommendations-*"))


def test_cancel_while_excel_is_being_compressed_cleans_everything(
    export, tmp_path, monkeypatch
):
    original = writer._PolledFile.write

    def stop(self, data):
        if not self.failed:
            self.failed = True
            raise ExportCancelled()
        return original(self, data)

    monkeypatch.setattr(writer._PolledFile, "write", stop)
    with pytest.raises(ExportCancelled):
        export([customer()])
    assert not (tmp_path / "Recommendations.zip").exists()
    assert not list(tmp_path.glob(".recommendations-*"))


def test_nonfinite_output_never_publishes_and_closes_xml(export, tmp_path):
    c = customer()
    p = SimpleNamespace(
        status=P.READY,
        probability=float("nan"),
        prediction_date="2026-09-01",
        target_until="2026-10-01",
        warnings=(),
    )
    with pytest.raises(ExportError, match="NONFINITE_EXPORT_NUMBER"):
        export([replace(c, propensity=p)])
    assert not (tmp_path / "Recommendations.zip").exists()
    assert not list(tmp_path.glob(".recommendations-*"))
