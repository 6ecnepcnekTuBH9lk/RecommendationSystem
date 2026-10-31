"""Exercise real subprocess success, failure, cancellation and deferred close."""

from dataclasses import asdict
from copy import deepcopy
from datetime import datetime
from pathlib import Path
import json
import os
import re
import time

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")

import pytest
from PyQt6.QtCore import QProcess, QTimer, Qt
from PyQt6.QtGui import QIcon
from PyQt6.QtTest import QTest
from PyQt6.QtWidgets import QApplication, QFrame, QLabel, QScrollArea, QTableWidget, QWidget, QVBoxLayout

from Application.mindbox.selection import DEFAULT_SELECTION
from Application.tabs import dataset_statistics_tab as ui
from Application.theme.apply_theme import apply_app_theme
from Application import statistics_cache as cache
from Application.settings.set_status import set_ready_status, schedule_status_reset

pytestmark = pytest.mark.usefixtures("window_settings")


def _visual_items(layout):
    """Preserve content-order checks across the explicit section containers."""
    items = []
    for index in range(layout.count()):
        item = layout.itemAt(index)
        section = item.layout()
        if section is not None and section.objectName() == 'statisticsSectionLayout':
            items.extend(section.itemAt(i) for i in range(section.count()))
        else:
            items.append(item)
    return items


@pytest.fixture
def app():
    application = QApplication.instance() or QApplication([])
    yield application
    application.setStyleSheet("")


@pytest.fixture
def sample_result():
    from Application.dataset_statistics import Coverage, DatasetStatistics

    stamp = "2026-09-24T09:47:28+00:00"
    interval = (("2025-01-01T00:00:00+00:00", "2026-01-01T00:00:00+00:00"),)
    return asdict(DatasetStatistics(
        calculated_at=stamp,
        coverage=tuple(Coverage(name, interval if name != "Customers" else (), ("API",), stamp)
                       for name in ("Actions", "Orders", "CustomerMerges", "Customers")),
        source_actions=3, total_interactions=3, orders=1, order_lines=1, customers=2, action_customers=2, order_customers=1,
        interaction_users=2, actions_with_product=2, actions_without_product=1,
        view_interactions=2, favorite_interactions=0, purchase_interactions=1, purchase_quantity="4",
        mean_interactions=1.5, median_interactions=1.5, unique_source_products=1, unique_resolved_items=1,
        resolved_interactions=3, unresolved_interactions=0, resolution_rate=100.,
        action_types=(("ProsmotrProduktaVApiMethod", 3),), line_statuses=(("CP", 1, True),),
        namespaces=(("offline1C", 3, 3, 0, 100.), ("unsupported", 0, 0, 0, 0.)),
        top_products=(("000001", "Рубашка", 2, 0, 1, 3),),
        diagnostics=tuple((key, int(key == "orders_unique")) for key in cache.DIAGNOSTIC_KEYS), warnings=(),
        gender_distribution=(("Мужчины", 0, 0.), ("Женщины", 0, 0.), ("Не указан", 2, 100.)),
        age_distribution=tuple((name, 2 if name == "Возраст не определён" else 0,
                                100. if name == "Возраст не определён" else 0.) for name in
                               ("До 18 лет", "18–25 лет", "26–35 лет", "36–45 лет", "46–55 лет",
                                "56–65 лет", "66 лет и старше", "Возраст не определён")),
        mean_age=None, median_age=None, view_users=1, favorite_users=0, purchase_users=1,
        view_purchase_users=0, favorite_purchase_users=0, all_interaction_type_users=0, repeat_buyers=0,
        interaction_activity_distribution=(("1", 1, 50.), ("2–5", 1, 50.), ("6–10", 0, 0.), ("11–25", 0, 0.),
                                           ("26–50", 0, 0.), ("51–100", 0, 0.), ("101+", 0, 0.)),
        purchase_order_distribution=(("1 заказ", 1, 100.), ("2 заказа", 0, 0.), ("3–5 заказов", 0, 0.),
                                     ("6–10 заказов", 0, 0.), ("11+ заказов", 0, 0.)),
        mean_orders_per_buyer=1., median_orders_per_buyer=1., repeat_buyer_rate=0., active_buyer_rate=50.,
        purchase_orders=1, orders_without_purchase=0, mean_purchase_lines_per_order=1., median_purchase_lines_per_order=1.,
        mean_purchase_units_per_order="4", median_purchase_units_per_order="4",
        purchase_basket_distribution=(("1 позиция", 1, 100.), ("2 позиции", 0, 0.), ("3–5 позиций", 0, 0.),
                                      ("6–10 позиций", 0, 0.), ("11+ позиций", 0, 0.)),
        order_financials=(("RUB", 1, 1, "4", "40", "40", "40"), ("KZT", 0, 0, "0", "0", "0", "0")),
        order_monthly_dynamics=(("2025-01", 1, "40", 0, "0"),),
        store_statistics=(("RUB", "shop", "Магазин", 1, 1, 1, "4", "40", "40"),),
        ordering_method_distribution=(("Не указано", 1, 100.),),
        delivery_type_distribution=(("Не указано", 1, 100.),),
        payment_type_distribution=(("Не указано", 1, 100.),),
        delivery_financials=(("RUB", 0, "0", "0", "0"), ("KZT", 0, "0", "0", "0")),
        mixed_currency_purchase_orders=0, unknown_currency_purchase_orders=0,
        mean_views_per_viewer=2., median_views_per_viewer=2., mean_favorites_per_user=0., median_favorites_per_user=0.,
        view_user_activity_distribution=tuple((label, int(label == "2–5"), 100. if label == "2–5" else 0.)
                                              for label in ("1", "2–5", "6–10", "11–25", "26–50", "51–100", "101+")),
        favorite_user_activity_distribution=tuple((label, 0, 0.) for label in ("1", "2", "3–5", "6–10", "11+")),
        action_channel_statistics=(("web", "Сайт", 2, 1, 100., 0, 0, 0.),),
        action_monthly_dynamics=(("2025-01", 2, 1, 0, 0),),
        view_parameter_actions=2,
        view_availability_distribution=(("Доступен", 1, 50.), ("Недоступен", 0, 0.), ("Не указано", 1, 50.)),
        view_price_statistics=(("RUB", 1, "100.25", "100.25"), ("KZT", 0, "0", "0")),
        products_with_views=1, products_with_favorites=0, products_with_purchases=1,
        resolved_view_interactions=2, resolved_favorite_interactions=0, resolved_purchase_interactions=1,
        resolved_purchase_quantity="4",
        top_viewed_products=(("000001", "Рубашка", 2, 1, 0, 1),), top_favorited_products=(),
        top_purchased_products=(("000001", "Рубашка", 1, 1, "4", 2, 0),),
        product_category_statistics=((None, "Не указано", 1, 2, 0, 1, "4"),),
        product_gender_statistics=(("Не указано", 1, 2, 0, 1, "4"),),
        product_season_statistics=(("Не указано", 1, 2, 0, 1, "4"),),
        product_style_statistics=(("Не указано", 1, 2, 0, 1, "4"),),
    ))


@pytest.fixture
def window(app):
    window = QWidget()
    window.status_label = QLabel(window)
    window.status_icon = QLabel(window)
    window._status_reset_timer = QTimer(window)
    window._status_reset_timer.setSingleShot(True)
    window._status_reset_timer.timeout.connect(lambda: set_ready_status(window))
    set_ready_status(window)
    layout = QVBoxLayout(window)
    tab = ui.DatasetStatisticsTab(window)
    layout.addWidget(tab)
    window.resize(1000, 700)
    window.show()
    yield window, tab
    tab._shutdown()
    window.close()
    window.deleteLater()
    app.processEvents()


def wait_until(predicate, timeout=10000):
    deadline = time.monotonic() + timeout / 1000
    while not predicate() and time.monotonic() < deadline:
        QTest.qWait(20)
    assert predicate()


def canonical_process(monkeypatch, tmp_path):
    (tmp_path / "canonical").mkdir()
    (tmp_path / "canonical/catalog.json").write_text(json.dumps({
        "schema_version": 2, "revision": "test", "actions": {}, "orders": {}, "customer_merges": None,
        "manual_interactions": None, "selection": asdict(DEFAULT_SELECTION)}))
    (tmp_path / "site_categories.csv").write_text("КодКатегории|НазваниеКатегории\n", encoding="utf-8-sig")
    catalog = tmp_path / "nomenclature.csv"
    catalog.write_text("КодНоменклатуры|Номенклатура\n000001|Test\n", encoding="utf-8")

    class Process(QProcess):
        def start(self, program, arguments):
            super().start(program, arguments + ["--raw-root", str(tmp_path), "--catalog", str(catalog),
                                               "--slice-cache", str(tmp_path / "slices")])

    monkeypatch.setattr(ui, "QProcess", Process)
    return catalog


def test_real_indexed_cli_in_existing_qprocess_then_error_preserves_display(window, monkeypatch, tmp_path):
    from tests.test_dataset_statistics import dataset as fixture
    from Application import statistics_index as index
    from Application.dataset_statistics import calculate_statistics
    _, tab = window
    dataset = fixture.__wrapped__(tmp_path / "fixture")
    root, reference, data = dataset
    monkeypatch.setattr(index, "protect_directory", lambda p: p.chmod(0o700))
    directory = tmp_path / "analytics"
    directory.mkdir()
    monkeypatch.setattr(index, "validate_storage_root", lambda p, **kw: index.path_identity(p))
    # This is a synthetic GUI/CLI fixture under pytest's shared workspace. ACL
    # policy has dedicated tests; give the child the same explicit fixture mock.
    cli = Path(ui.__file__).resolve().parents[2]/"scripts/dataset_statistics.py"
    bootstrap = tmp_path/"synthetic_statistics_cli.py"
    bootstrap.write_text(
        "import sys, runpy\n"
        f"sys.path.insert(0, {str(cli.parent.parent)!r})\n"
        "from Application import statistics_index as index\n"
        "index.validate_storage_root=lambda p, **kw: index.path_identity(p)\n"
        f"runpy.run_path({str(cli)!r}, run_name='__main__')\n", encoding="utf-8")
    # Build on an earlier controlled local day; the actual child QProcess must
    # still use INDEXED today, computing today's demographics from the snapshot.
    original_capture = index.capture_sources

    def earlier_capture(*args, **kwargs):
        return original_capture(*args, **kwargs, calculated_at=datetime.fromisoformat("2000-01-01T12:00:00+00:00"))

    with monkeypatch.context() as context:
        context.setattr(index, "capture_sources", earlier_capture)
        index.build_statistics_index(raw_root=root, catalog_path=reference, index_dir=directory, expected_bytes=1024**2)
    expected = asdict(calculate_statistics(raw_root=root, catalog_path=reference))

    class Process(QProcess):
        def start(self, program, arguments):
            arguments = [str(bootstrap) if Path(a).name == "dataset_statistics.py" else a for a in arguments]
            super().start(program, arguments + ["--raw-root", str(root), "--catalog", str(reference),
                "--slice-cache", str(tmp_path / "slices"), "--index-dir", str(directory), "--backend", "INDEXED"])

    monkeypatch.setattr(ui, "QProcess", Process)
    tab.start()
    wait_until(lambda: tab.process is None)
    assert tab.failure is None
    displayed = cache.load_result(tab.cache_path)
    assert displayed is not None
    actual = deepcopy(displayed)
    actual.pop("calculated_at")
    expected.pop("calculated_at")
    assert actual == json.loads(json.dumps(expected))
    assert tab.export_button.isEnabled()
    data["revision"] = "updated"
    (root / "canonical/catalog.json").write_text(json.dumps(data))
    tab.start()
    wait_until(lambda: tab.process is None)
    assert tab.failure
    assert cache.load_result(tab.cache_path) == displayed
    assert tab.export_button.isEnabled()


def test_success_theme_switch_and_error_preserves_previous_result(window, monkeypatch, tmp_path):
    _, tab = window
    catalog = canonical_process(monkeypatch, tmp_path)
    tab.start()
    process = tab.process
    assert not tab.export_button.isEnabled()
    tab.start()
    assert tab.process is process
    wait_until(lambda: tab.process is None)
    assert tab.export_button.isEnabled()
    assert tab.card_labels[0].text() == "0"
    assert tab.card_labels[3].text() == "—"
    assert tab.sections.count() == 5
    saved = tab.cache_path.read_bytes()
    assert cache.load_result(tab.cache_path)["total_interactions"] == 0
    from Application.dataset_statistics import calculate_statistics

    result = asdict(calculate_statistics(raw_root=tmp_path, catalog_path=catalog))
    result["resolution_rate"] = 99.99834283990839
    tab.render(result)
    table = tab.sections.widget(4).findChildren(QTableWidget)[1]
    assert table.item(table.rowCount() - 1, 1).text() == "99.9983"
    for dark in (False, True):
        apply_app_theme(QApplication.instance(), dark)
        QApplication.processEvents()
        assert tab.card_labels[0].font().pointSize() == 22
        assert tab.card_labels[0].text() == "0"
    catalog.unlink()
    tab.start()
    wait_until(lambda: tab.process is None)
    assert "Catalog" in tab.status.text()
    assert "не обновлены" in tab.status.text()
    assert tab.card_labels[0].text() == "0"
    assert tab.refresh.isEnabled() and not tab.cancel_button.isEnabled()
    assert tab.export_button.isEnabled()
    assert tab.cache_path.read_bytes() == saved
    wait_until(lambda: tab.window.status_label.text() == "Готов к работе")


@pytest.mark.parametrize("close", [False, True])
def test_cancel_and_close_reap_child_without_blocking_event_loop(window, monkeypatch, close, sample_result):
    window, tab = window
    cache.save_result(tab.cache_path, sample_result)
    tab.render(sample_result)
    saved = tab.cache_path.read_bytes()
    assert tab.export_button.isEnabled()

    class SlowProcess(QProcess):
        def start(self, program, arguments):
            super().start(program, ["-c", "import time; time.sleep(30)"])

    monkeypatch.setattr(ui, "QProcess", SlowProcess)
    ticks = []
    timer = QTimer(tab)
    timer.timeout.connect(lambda: ticks.append(1))
    timer.start(10)
    tab.start()
    wait_until(lambda: tab.process.state() == QProcess.ProcessState.Running)
    assert not tab.export_button.isEnabled()
    QTest.qWait(50)
    if close:
        window.close()
    else:
        tab.cancel()
    wait_until(lambda: tab.process is None)
    wait_until(lambda: not window.isVisible() if close else tab.refresh.isEnabled())
    assert ticks
    assert "отменён" in tab.status.text()
    assert not tab.kill_timer.isActive()
    assert tab.cache_path.read_bytes() == saved
    assert tab.card_labels[0].text() == "3"
    assert tab.export_button.isEnabled()
    wait_until(lambda: window.status_label.text() == "Готов к работе")


def test_failed_start_resets_controls(window, monkeypatch):
    _, tab = window
    monkeypatch.setattr(ui.sys, "executable", "missing-statistics-python-executable")
    tab.start()
    wait_until(lambda: tab.process is None)
    assert "Не удалось запустить" in tab.status.text()
    assert tab.refresh.isEnabled()
    assert not tab.cancel_button.isEnabled()
    assert not tab.export_button.isEnabled()
    wait_until(lambda: tab.window.status_label.text() == "Готов к работе")


def test_initial_page_has_no_internal_heading_or_cache(window):
    _, tab = window
    assert tab.layout().itemAt(0).widget() is tab.description
    assert tab.description.text() == "Исходные данные Mindbox. Отбор не установлен."
    assert tab.status.text() == "В фоновом режиме будут рассчитаны данные из Mindbox."
    assert not [label for label in tab.findChildren(QLabel) if label.property("class") == "statisticsSection"]
    assert [label.text() for label in tab.card_labels] == ["—"] * 4
    assert not tab.cache_path.exists()
    assert card_titles(tab)[:4] == ["Количество взаимодействий", "Количество заказов",
                                   "Количество позиций в заказах", "Профили текущего среза"]
    controls = tab.layout().itemAt(2).layout()
    assert controls.count() == 6
    assert [controls.itemAt(i).widget() for i in range(6)] == [
        tab.refresh, tab.export_button, tab.detail_export_button, tab.cancel_button, tab.status, tab.progress_container]
    assert controls.stretch(5) == 1
    assert not tab.export_button.isEnabled()


def card_titles(page):
    return [card.findChildren(QLabel)[0].text() for card in page.findChildren(QFrame)
            if card.property("class") == "statisticsCard"]


def assert_section_names(tab):
    assert [tab.sections.tabText(i) for i in range(tab.sections.count())] == [
        "Действия", "Заказы", "Товары", "Клиенты", "Техническая информация"]


@pytest.mark.parametrize("outside,fractional", [(0, 0), (7, 13)])
def test_technical_information_renders_order_filter_diagnostics(window, sample_result, outside, fractional):
    _, tab = window
    values = dict(sample_result["diagnostics"])
    values.update(orders_outside_statistics_period=outside, fractional_quantity_order_lines=fractional)
    sample_result["diagnostics"] = tuple(values.items())
    tab.render(sample_result)
    tab.sections.setCurrentIndex(4)
    QApplication.processEvents()
    assert tab.sections.tabText(4) == "Техническая информация"
    diagnostics = next(table for table in tab.sections.widget(4).findChildren(QTableWidget)
                       if table.horizontalHeaderItem(0).text() == "Диагностика")
    rows = {diagnostics.item(row, 0).text(): diagnostics.item(row, 1).text()
            for row in range(diagnostics.rowCount())}
    assert rows["Заказы вне периода статистики"] == str(outside)
    assert rows["Позиции с нецелым количеством"] == str(fractional)
    assert not diagnostics.isHidden()


def test_five_pages_technical_content_scroll_and_no_tooltips(window, sample_result):
    _, tab = window
    assert_section_names(tab)
    for _ in range(2):
        tab.render(sample_result)
        assert_section_names(tab)
        technical = tab.sections.widget(4)
        assert isinstance(technical, QScrollArea)
        assert technical.widgetResizable()
        layout = technical.widget().layout()
        assert layout.spacing() == ui.SECTION_SPACING
        assert layout.count() == 5
        coverage, summary, raw, quality, diagnostics = technical.findChildren(QTableWidget)
        assert technical.findChildren(QTableWidget) == [coverage, summary, raw, quality, diagnostics]
        assert [coverage.horizontalHeaderItem(i).text() for i in range(3)] == [
            "Источник", "Период / состояние", "Источник данных"]
        assert [coverage.item(i, 0).text() for i in range(4)] == [
            "Действия", "Заказы", "Объединения клиентов", "Клиенты"]
        assert summary.rowCount() == 13
        assert summary.item(0, 1).text() == "3"
        assert summary.item(12, 1).text() == "100.0000"
        assert diagnostics.horizontalHeaderItem(0).text() == "Диагностика"
        assert diagnostics.rowCount() == len(sample_result["diagnostics"]) + 5
        assert [diagnostics.item(i, 0).text() for i in range(len(sample_result["diagnostics"]))] == [
            ui.DIAGNOSTIC_LABELS[key] for key, _ in sample_result["diagnostics"]]
        for index, counts in enumerate((8, 7, 9, 12)):
            assert len(tab.sections.widget(index).findChildren(QTableWidget)) == counts
        customers = next(table for table in tab.sections.widget(3).findChildren(QTableWidget)
                         if table.horizontalHeaderItem(0).text() == 'Пол')
        assert customers.rowCount() == 3
        assert customers.item(2, 1).text() == "2"
        for index in range(5):
            for table in tab.sections.widget(index).findChildren(QTableWidget):
                for row in range(table.rowCount()):
                    for column in range(table.columnCount()):
                        assert table.item(row, column).toolTip() == ""
    tab.sections.setCurrentIndex(4)
    QApplication.processEvents()
    scrollbar = technical.verticalScrollBar()
    assert scrollbar.maximum() > 0
    blocks = [coverage, summary, diagnostics]
    assert all(a.geometry().bottom() < b.geometry().top() for a, b in zip(blocks, blocks[1:]))
    scrollbar.setValue(scrollbar.maximum())
    QApplication.processEvents()
    # The last table's bottom is reachable, independent of theme/font pixel sizes.
    bottom = diagnostics.mapTo(technical.viewport(), diagnostics.rect().bottomLeft())
    assert technical.viewport().rect().contains(bottom)


def test_russian_presentation_dates_spacing_and_original_values(window, sample_result):
    _, tab = window
    original = deepcopy(sample_result)
    sample_result["warnings"] = tuple(ui.WARNING_LABELS)
    tab.render(sample_result)
    expected = datetime.fromisoformat(sample_result["calculated_at"]).astimezone().strftime("%d.%m.%Y %H:%M:%S")
    assert tab.description.text() == "Исходные данные Mindbox. Отбор не установлен. Дата и время расчета: " + expected
    assert re.fullmatch(r"Исходные данные Mindbox\. Отбор не установлен\. Дата и время расчета: "
                        r"\d{2}\.\d{2}\.\d{4} \d{2}:\d{2}:\d{2}", tab.description.text())
    assert tab.status.text().splitlines()[0] == 'Статистика рассчитана за период: 01.01.2025 - 01.01.2026.'
    coverage = tab.sections.widget(4).findChildren(QTableWidget)[0]
    assert coverage.item(0, 1).text() == "01.01.2025 — 01.01.2026"
    assert coverage.item(3, 1).text().startswith("Снимок данных от ")
    assert card_titles(tab.sections.widget(0)) == ["Просмотры товаров", "Добавления в избранное", "Клиенты с просмотрами", "Клиенты с избранным",
                                                "Доля VIEW/FAVORITE без товара, %"]
    assert card_titles(tab.sections.widget(1))[:4] == ["Заказы с покупкой", "Невыкупленные заказы", "Позиции покупок", "Покупатели"]
    assert card_titles(tab.sections.widget(2)) == ["Товары с взаимодействиями", "Товары с просмотрами",
                                                "Товары в избранном", "Купленные товары"]
    for index in range(5):
        assert tab.sections.widget(index).widget().layout().spacing() == ui.SECTION_SPACING
    labels = [label.text() for label in tab.findChildren(QLabel)]
    cells = []
    for table in tab.findChildren(QTableWidget):
        labels.extend(table.horizontalHeaderItem(i).text() for i in range(table.columnCount()))
        cells.extend(table.item(r, c).text() for r in range(table.rowCount()) for c in range(table.columnCount()))
    for value in ("ProsmotrProduktaVApiMethod", "offline1C", "CP"):
        assert value in cells  # Technical source values are deliberately preserved.
    technical_values = {"ProsmotrProduktaVApiMethod", "offline1C", "CP"}
    presentation = "\n".join(labels + [value for value in cells if value not in technical_values])
    for phrase in ("Canonical", "canonical", "Snapshot", "item interactions",
                   "Raw events", "System name", "PURCHASE", "Resolution", "Namespace",
                   "Unique", "Resolved", "Unresolved", "Active users", "unsupported"):
        assert phrase not in presentation
    assert sample_result["coverage"] == original["coverage"]
    assert sample_result["purchase_quantity"] == "4"
    assert sample_result["resolution_rate"] == 100.


def test_conflicting_order_warning_is_hidden_only_from_top_status(window, sample_result):
    _, tab = window
    warning = 'Есть конфликтующие снимки заказов: item interactions рассчитаны по первому снимку.'
    sample_result['warnings'] = (warning,)
    sample_result['diagnostics'] = (*sample_result['diagnostics'], ('orders_duplicate_conflicting', 7))
    original = deepcopy(sample_result)
    tab.render(sample_result)
    assert tab.status.text() == 'Статистика рассчитана за период: 01.01.2025 - 01.01.2026.'
    assert sample_result == original and tab.displayed_result['warnings'] == (warning,)
    tables = tab.sections.widget(4).findChildren(QTableWidget)
    assert any(table.item(row, 0).text() == ui.DIAGNOSTIC_LABELS['orders_duplicate_conflicting']
               and table.item(row, 1).text() == '7' for table in tables for row in range(table.rowCount()))


def test_saved_result_loads_automatically_without_process(app, sample_result, monkeypatch):
    cache.save_result(ui.CACHE_PATH, sample_result)
    monkeypatch.setattr(ui.DatasetStatisticsTab, "start", lambda *args: pytest.fail("Must not recalculate"))
    window = QWidget()
    tab = ui.DatasetStatisticsTab(window)
    try:
        assert tab.process is None
        assert tab.card_labels[0].text() == "3"
        expected = datetime.fromisoformat(sample_result["calculated_at"]).astimezone().strftime("%d.%m.%Y %H:%M:%S")
        assert tab.description.text() == "Исходные данные Mindbox. Отбор не установлен. Дата и время расчета: " + expected
        assert tab.status.text() == 'Статистика рассчитана за период: 01.01.2025 - 01.01.2026.'
        assert_section_names(tab)
        assert len(tab.sections.widget(4).findChildren(QTableWidget)) == 5
    finally:
        window.deleteLater()
        app.processEvents()


@pytest.mark.parametrize("content", [b'{"schema_version":', b'{}', b'[]', b'\xff\xfe',
                                     b'{"schema_version":1,"result":{}}'])
def test_corrupt_cache_starts_empty_without_overwriting_file(app, content, caplog):
    ui.CACHE_PATH.write_bytes(content)
    window = QWidget()
    tab = ui.DatasetStatisticsTab(window)
    try:
        assert tab.description.text() == "Исходные данные Mindbox. Отбор не установлен."
        assert [label.text() for label in tab.card_labels] == ["—"] * 4
        assert ui.CACHE_PATH.read_bytes() == content
        assert "Не удалось загрузить" in caplog.text
        assert "Traceback" not in tab.status.text()
    finally:
        window.deleteLater()
        app.processEvents()


def result_process(monkeypatch, result, *, suffix="", exit_code=0):
    message = json.dumps({"event": "result", "value": result}) + "\n"

    class Process(QProcess):
        def start(self, program, arguments):
            script = f"import sys; sys.stdout.write({(message + suffix)!r}); sys.stdout.flush(); sys.exit({exit_code})"
            super().start(program, ["-c", script])

    monkeypatch.setattr(ui, "QProcess", Process)


def set_profile_count(result, count):
    result["customers"] = count
    for key in ("gender_distribution", "age_distribution"):
        result[key] = tuple((label, count if i == len(result[key]) - 1 else 0,
                             100. if i == len(result[key]) - 1 else 0.)
                            for i, (label, _, _) in enumerate(result[key]))


def test_new_success_atomically_replaces_cache_before_render(window, sample_result, monkeypatch):
    _, tab = window
    cache.save_result(tab.cache_path, sample_result)
    previous = tab.cache_path.read_bytes()
    candidate = deepcopy(sample_result)
    candidate["calculated_at"] = "2026-09-25T09:47:28+00:00"
    set_profile_count(candidate, 5)
    result_process(monkeypatch, candidate)
    render = tab.render

    def checked_render(result):
        saved = cache.load_result(tab.cache_path)
        assert saved["calculated_at"] == candidate["calculated_at"]
        assert saved["customers"] == 5
        render(result)

    monkeypatch.setattr(tab, "render", checked_render)
    tab.start()
    wait_until(lambda: tab.process is None)
    assert tab.card_labels[3].text() == "5"
    assert tab.cache_path.read_bytes() != previous
    assert not list(tab.cache_path.parent.glob(".dataset-statistics-*.tmp"))
    assert [p.name for p in tab.cache_path.parent.glob("*statistics*.json")] == [tab.cache_path.name]
    assert tab.status.text() == 'Статистика рассчитана за период: 01.01.2025 - 01.01.2026.'
    wait_until(lambda: tab.window.status_label.text() == "Готов к работе")


@pytest.mark.parametrize("failure", ["incomplete", "bad_type", "bad_date", "bad_rows", "nan", "trailing_json",
                                    "json_error", "deep_json", "prior_incomplete", "nonzero_exit", "crash", "replace", "fsync", "order_money", "action_price", "product_quantity", "orders_without_purchase"])
def test_failed_new_result_preserves_cache_and_display(window, sample_result, monkeypatch, failure):
    _, tab = window
    cache.save_result(tab.cache_path, sample_result)
    tab.render(sample_result)
    previous = tab.cache_path.read_bytes()
    candidate = deepcopy(sample_result)
    set_profile_count(candidate, 9)
    if failure == "incomplete":
        del candidate["top_products"]
    elif failure == "bad_type":
        candidate["total_interactions"] = True
    elif failure == "bad_date":
        candidate["calculated_at"] = "not a date"
    elif failure == "bad_rows":
        candidate["namespaces"] = [["offline1C"]]
    elif failure == "nan":
        candidate["resolution_rate"] = float("nan")
    elif failure == "orders_without_purchase":
        candidate["orders_without_purchase"] = 1
    elif failure == "product_quantity":
        candidate["resolved_purchase_quantity"] = "1.5"
    elif failure == "action_price":
        candidate["view_price_statistics"] = (("RUB", 1, "NaN", "0"), ("KZT", 0, "0", "0"))
    elif failure == "order_money":
        candidate["mean_purchase_units_per_order"] = "NaN"
    elif failure in ("replace", "fsync"):
        def fail(*args):
            raise OSError("Нет доступа к файлу")
        monkeypatch.setattr(cache.os, failure, fail)
    result_process(monkeypatch, candidate,
                   suffix={"trailing_json": "{", "json_error": "bad JSON\n",
                           "deep_json": "[" * 1200 + "0" + "]" * 1200 + "\n"}.get(failure, ""),
                   exit_code=1 if failure == "nonzero_exit" else 0)
    if failure == "prior_incomplete":
        result_process(monkeypatch, None, suffix=json.dumps({"event": "result", "value": candidate}) + "\n")
    if failure == "crash":
        class CrashProcess(QProcess):
            def start(self, program, arguments):
                super().start(program, ["-c", "import time; time.sleep(30)"])
        monkeypatch.setattr(ui, "QProcess", CrashProcess)
    tab.start()
    if failure == "crash":
        wait_until(lambda: tab.process.state() == QProcess.ProcessState.Running)
        tab.process.kill()
    wait_until(lambda: tab.process is None)
    assert tab.cache_path.read_bytes() == previous
    assert tab.card_labels[3].text() == "2"
    assert "не обновлены" in tab.status.text()
    assert tab.refresh.isEnabled() and not tab.cancel_button.isEnabled()
    assert not list(tab.cache_path.parent.glob(".dataset-statistics-*.tmp"))
    wait_until(lambda: tab.window.status_label.text() == "Готов к работе")


def test_information_style_and_existing_icons_in_both_themes(window):
    _, tab = window
    for dark in (False, True):
        apply_app_theme(QApplication.instance(), dark)
        QApplication.processEvents()
        for label in (tab.description, tab.status):
            assert label.property("class") == "statisticsInfo"
            alignment = Qt.AlignmentFlag.AlignHCenter if label is tab.description else Qt.AlignmentFlag.AlignLeft
            assert label.alignment() & alignment
            assert label.font().italic()
            assert label.font().pointSize() == 10
        for button, filename in ((tab.refresh, "statistic.png"), (tab.export_button, "export.png"),
                                 (tab.cancel_button, "failure.png")):
            expected = QIcon(str(ui.ICONS_DIR / filename))
            assert not button.icon().isNull()
            assert button.icon().pixmap(button.iconSize()).toImage() == expected.pixmap(button.iconSize()).toImage()


def test_progress_and_shared_status_resist_stale_reset(window, monkeypatch):
    window, tab = window
    schedule_status_reset(window, 5)
    window._status_reset_timer.start(20)

    class ProgressProcess(QProcess):
        def start(self, program, arguments):
            message = json.dumps({"event": "progress", "value": "actions: 10 000"})
            super().start(program, ["-u", "-c", f"import time; print({message!r}, flush=True); time.sleep(30)"])

    monkeypatch.setattr(ui, "QProcess", ProgressProcess)
    tab.start()
    assert window.status_label.text() == "Расчет статистики..."
    assert not window._status_reset_timer.isActive()
    wait_until(lambda: tab.status.text() == "Расчет → Действия: 10 000")
    QTest.qWait(50)
    assert tab.process is not None
    assert window.status_label.text() == "Расчет статистики..."
    # Even a later reset requested elsewhere cannot leave this process "ready".
    schedule_status_reset(window, 0)
    QTest.qWait(30)
    assert window.status_label.text() == "Расчет статистики..."
    tab.cancel()
    wait_until(lambda: tab.process is None)
    wait_until(lambda: window.status_label.text() == "Готов к работе")


@pytest.mark.parametrize("actions,orders,expected", [
    ([("2024-01-01", "2024-12-01")], [("2024-06-01", "2025-01-01")],
     'Статистика рассчитана за период: 01.06.2024 - 01.12.2024.'),
    ([("2024-01-01", "2024-02-01"), ("2024-03-01", "2024-04-01")], [("2024-01-01", "2024-04-01")],
     'Статистика рассчитана за периоды: 01.01.2024 - 01.02.2024; 01.03.2024 - 01.04.2024.'),
    ([("2024-01-01", "2024-02-01")], [("2024-02-01", "2024-03-01")], "Статистика рассчитана."),
    ([], [("2024-06-01", "2024-07-01")], 'Статистика рассчитана за период: 01.06.2024 - 01.07.2024.'),
    ([], [], "Статистика рассчитана."),
])
def test_period_uses_interaction_coverage_not_customer_snapshots(sample_result, actions, orders, expected):
    for source, intervals in zip(sample_result["coverage"][:2], (actions, orders)):
        source["intervals"] = [(a + "T00:00:00+00:00", b + "T00:00:00+00:00") for a, b in intervals]
    # Customers and merges still cover 2025: they must not enter the calculation.
    sample_result["coverage"][3]["intervals"] = sample_result["coverage"][2]["intervals"]
    assert ui._calculation_message(sample_result) == expected


def test_customer_analytics_render_themes_scroll_and_no_duplication(window, sample_result):
    _, tab = window
    for dark in (False, True):
        apply_app_theme(QApplication.instance(), dark)
        tab.render(sample_result)
        tab.sections.setCurrentIndex(3)
        QApplication.processEvents()
        page = tab.sections.widget(3)
        assert page.widgetResizable()
        assert page.widget().layout().spacing() == ui.SECTION_SPACING
        assert card_titles(page) == ["Профили текущего среза", "Клиенты с товарными взаимодействиями", "Покупатели", "Повторные покупатели (от 2 покупок)",
                                    "Средний возраст", "Медианный возраст"]
        assert [label.text() for label in page.findChildren(QLabel) if label.property("class") == "statisticsNumber"] == [
            "2", "2", "1", "0", "—", "—"]
        headings = [label.text() for label in page.findChildren(QLabel) if label.property("class") == "statisticsSection"]
        assert headings == ["ОХВАТ КЛИЕНТСКОЙ БАЗЫ", "ЦИФРОВАЯ АКТИВНОСТЬ КЛИЕНТОВ", "ПОКУПАТЕЛИ И КЛИЕНТЫ БЕЗ ПОКУПОК",
                            "ПОРТРЕТ КЛИЕНТА", "ВОЗРАСТНАЯ СТРУКТУРА КЛИЕНТОВ", "РАСПРЕДЕЛЕНИЕ ПО ПОЛУ",
                            "АКТИВНОСТЬ КЛИЕНТОВ", "РАСПРЕДЕЛЕНИЕ КЛИЕНТОВ ПО КОЛИЧЕСТВУ ВЗАИМОДЕЙСТВИЙ", "ПОКУПАТЕЛЬСКАЯ АКТИВНОСТЬ",
                            "ПОКУПАТЕЛЬСКАЯ АНАЛИТИКА — STRICT / EXTENDED", "RFM-АНАЛИЗ ПОКУПАТЕЛЕЙ", "КОГОРТНЫЙ АНАЛИЗ ПОКУПАТЕЛЕЙ"]
        tables = page.findChildren(QTableWidget)
        assert [table.rowCount() for table in tables] == [7, 5, 4, 8, 3, 9, 7, 5, 3, 1, 1, 1]
        assert tables[5].item(8, 1).text() == "50.00"
        for table in tables:
            assert all(table.item(r, c).toolTip() == "" for r in range(table.rowCount()) for c in range(table.columnCount()))
        assert page.verticalScrollBar().maximum() > 0
        page.verticalScrollBar().setValue(page.verticalScrollBar().maximum())
        QApplication.processEvents()
        bottom = tables[-1].mapTo(page.viewport(), tables[-1].rect().bottomLeft())
        assert page.viewport().rect().contains(bottom)
        assert_section_names(tab)


@pytest.mark.parametrize('customers', [2, None])
def test_customer_demographic_layout_order_and_missing_profiles(window, sample_result, customers):
    from copy import deepcopy
    from Application.statistics_charts import StatisticsChart
    from test_statistics_charts import _populated
    from PyQt6.QtCore import QCoreApplication, QEvent
    _, tab = window
    result = _populated(sample_result)
    result['customers'] = customers
    before = deepcopy(result)
    for _ in range(2):
        tab.render(result)
        QCoreApplication.sendPostedEvents(None, QEvent.Type.DeferredDelete)
        QApplication.processEvents()
        page = tab.sections.widget(3)
        layout = page.widget().layout()
        items = _visual_items(layout)
        assert items[0].layout().count() == 4
        assert items[1].widget().text() == 'ОХВАТ КЛИЕНТСКОЙ БАЗЫ'
        profile_index = next(i for i, item in enumerate(items)
                             if isinstance(item.widget(), QLabel) and item.widget().text() == 'ПОРТРЕТ КЛИЕНТА')
        index = profile_index + 1
        if customers is None:
            assert items[index].widget().text() == 'Данные профилей клиентов отсутствуют. Пол и возраст недоступны.'
            index += 1
        assert items[index].widget().text() == 'ВОЗРАСТНАЯ СТРУКТУРА КЛИЕНТОВ'
        cards = items[index + 1].layout()
        assert cards.count() == 2
        assert [cards.itemAt(i).widget().findChildren(QLabel)[0].text() for i in range(2)] == ['Средний возраст', 'Медианный возраст']
        age = items[index + 2].widget()
        age_table = items[index + 3].widget()
        assert isinstance(age, StatisticsChart) and age.data.title == 'Возрастная структура клиентов'
        assert isinstance(age_table, QTableWidget) and age_table.horizontalHeaderItem(0).text() == 'Возрастная группа'
        assert items[index + 4].widget().text() == 'РАСПРЕДЕЛЕНИЕ ПО ПОЛУ'
        gender = items[index + 5].widget()
        gender_table = items[index + 6].widget()
        assert isinstance(gender, StatisticsChart) and gender.data.kind == 'donut'
        assert isinstance(gender_table, QTableWidget) and gender_table.horizontalHeaderItem(0).text() == 'Пол'
        assert items[index + 7].widget().text() == 'АКТИВНОСТЬ КЛИЕНТОВ'
        assert len(page.findChildren(QTableWidget)) == 12
        assert len(page.findChildren(StatisticsChart)) == 3
        assert len(tab.findChildren(StatisticsChart)) == 17
        assert card_titles(page) == ['Профили текущего среза', 'Клиенты с товарными взаимодействиями', 'Покупатели',
                                    'Повторные покупатели (от 2 покупок)', 'Средний возраст', 'Медианный возраст']
        assert result == before


def test_old_complete_cache_opens_empty_without_modifying_it(app, sample_result):
    # Even a complete payload with the old envelope must not be interpreted as v9.
    ui.CACHE_PATH.write_text(json.dumps({"schema_version": 8, "result": sample_result}), encoding="utf-8")
    saved = ui.CACHE_PATH.read_bytes()
    window = QWidget()
    tab = ui.DatasetStatisticsTab(window)
    try:
        assert [label.text() for label in tab.card_labels] == ["—"] * 4
        assert not tab.sections.widget(3).findChildren(QTableWidget)
        assert ui.CACHE_PATH.read_bytes() == saved
        assert_section_names(tab)
    finally:
        window.deleteLater()
        app.processEvents()


def test_orders_analytics_headers_themes_scroll_and_immutable_display(window, sample_result):
    _, tab = window
    original = deepcopy(sample_result)
    expected_headers = [
        ["Количество позиций", "Заказов", "Доля, %"],
        ["Валюта", "Заказов", "Позиции", "Единиц", "Сумма покупок", "Средняя сумма заказа", "Медианная сумма заказа"],
        ["Месяц", "Заказы RUB", "Сумма RUB", "Заказы KZT", "Сумма KZT"],
        ["Магазин / канал", "Заказов", "Покупателей", "Позиции", "Единиц", "Сумма покупок", "Средняя сумма заказа"],
        ["Магазин / канал", "Заказов", "Покупателей", "Позиции", "Единиц", "Сумма покупок", "Средняя сумма заказа"],
        ["Способ оплаты", "Заказов с этим способом", "Доля заказов, %"],
        ["Статус", "Позиции", "Доля, %", "Считается покупкой"],
    ]
    for dark in (False, True):
        apply_app_theme(QApplication.instance(), dark)
        tab.render(sample_result)
        assert_section_names(tab)
        tab.sections.setCurrentIndex(1)
        QApplication.processEvents()
        page = tab.sections.widget(1)
        assert page.widgetResizable() and page.widget().layout().spacing() == ui.SECTION_SPACING
        assert card_titles(page)[:4] == ["Заказы с покупкой", "Невыкупленные заказы", "Позиции покупок", "Покупатели"]
        assert len(card_titles(page)) == 8
        labels = [label.text() for label in page.findChildren(QLabel)]
        assert [label.text() for label in page.findChildren(QLabel) if label.property("class") == "statisticsSection"] == [
            "КОРЗИНА ЗАКАЗА", "ФИНАНСОВЫЕ ПОКАЗАТЕЛИ", "ДИНАМИКА КОЛИЧЕСТВА ПОКУПОК — RUB", "ДИНАМИКА КОЛИЧЕСТВА ПОКУПОК — KZT", "ДИНАМИКА СУММЫ ПОКУПОК — RUB", "ДИНАМИКА СУММЫ ПОКУПОК — KZT",
            "МАГАЗИНЫ И КАНАЛЫ — RUB",
            "МАГАЗИНЫ И КАНАЛЫ — KZT", "ОПЛАТА", "СТАТУСЫ ПОЗИЦИЙ"]
        assert not any("превышать 100%" in label for label in labels)
        assert not any(label.startswith("RUB —") for label in labels)
        assert not any(label.startswith("KZT —") for label in labels)
        tables = page.findChildren(QTableWidget)
        assert [[table.horizontalHeaderItem(i).text() for i in range(table.columnCount())] for table in tables] == expected_headers
        assert [tables[1].item(i, 0).text() for i in range(2)] == ["RUB", "KZT"]
        assert tables[1].rowCount() == 2  # No cross-currency total.
        assert tables[1].item(0, 4).text() == "40.00"
        assert tables[2].item(0, 0).text() == "01.2025"
        assert _visual_items(page.widget().layout())[-1].widget() is tables[-1]
        for table in tables:
            assert all(table.item(r, c).toolTip() == "" for r in range(table.rowCount()) for c in range(table.columnCount()))
        assert page.verticalScrollBar().maximum() > 0
        page.verticalScrollBar().setValue(page.verticalScrollBar().maximum())
        QApplication.processEvents()
        assert page.viewport().rect().contains(tables[-1].mapTo(page.viewport(), tables[-1].rect().bottomLeft()))
    assert original == sample_result
    assert ui._decimal_display("1234567.125") == "1 234 567.12"
    assert ui._decimal_display("1234.500", money=False) == "1 234.5"



def test_action_analytics_headers_themes_scroll_and_immutable_display(window, sample_result):
    _, tab = window
    original = deepcopy(sample_result)
    headers = [
        ["Показатель", "Значение"],
        ["Тип действия", "Количество взаимодействий", "Количество клиентов", "Среднее на клиента", "Медиана на клиента"],
        ["Количество просмотров", "Количество клиентов", "Доля, %"],
        ["Количество избранного", "Количество клиентов", "Доля, %"],
        ["Канал", "Просмотры", "Клиенты с просмотрами", "Доля просмотров, %", "Избранное", "Клиенты с избранным", "Доля избранного, %"],
        ["Месяц", "Просмотры", "Клиенты с просмотрами", "Избранное", "Клиенты с избранным"],
        ["Доступность", "Просмотров", "Доля, %"],
        ["Валюта", "Просмотров с ценой", "Средняя цена", "Медианная цена"],
    ]
    for dark in (False, True):
        apply_app_theme(QApplication.instance(), dark)
        for _ in range(2):
            tab.render(sample_result)
            assert_section_names(tab)
            tab.sections.setCurrentIndex(0)
            QApplication.processEvents()
            page = tab.sections.widget(0)
            assert card_titles(page) == ["Просмотры товаров", "Добавления в избранное", "Клиенты с просмотрами", "Клиенты с избранным",
                                         "Доля VIEW/FAVORITE без товара, %"]
            assert [label.text() for label in page.findChildren(QLabel) if label.property("class") == "statisticsNumber"] == ["2", "0", "1", "0", "Нет данных"]
            assert [label.text() for label in page.findChildren(QLabel) if label.property("class") == "statisticsSection"] == [
                "КАЧЕСТВО ИСХОДНЫХ СОБЫТИЙ",
                "АКТИВНОСТЬ КЛИЕНТОВ", "РАСПРЕДЕЛЕНИЕ КЛИЕНТОВ ПО КОЛИЧЕСТВУ ПРОСМОТРОВ",
                "РАСПРЕДЕЛЕНИЕ КЛИЕНТОВ ПО КОЛИЧЕСТВУ ИЗБРАННОГО", "КАНАЛЫ",
                "ДИНАМИКА ДЕЙСТВИЙ", "ПАРАМЕТРЫ ПРОСМОТРОВ"]
            assert page.widgetResizable() and page.widget().layout().spacing() == ui.SECTION_SPACING
            tables = page.findChildren(QTableWidget)
            assert [[table.horizontalHeaderItem(i).text() for i in range(table.columnCount())] for table in tables] == headers
            assert tables[5].item(0, 0).text() == "01.2025"
            assert tables[7].rowCount() == 2
            assert [tables[7].item(i, 0).text() for i in range(2)] == ["RUB", "KZT"]
            assert tables[7].item(0, 2).text() == "100.25"
            for table in tables:
                assert all(table.item(r, c).toolTip() == "" for r in range(table.rowCount()) for c in range(table.columnCount()))
            wait_until(lambda: page.verticalScrollBar().maximum() > 0)
            page.verticalScrollBar().setValue(page.verticalScrollBar().maximum())
            QApplication.processEvents()
            assert page.viewport().rect().contains(tables[-1].mapTo(page.viewport(), tables[-1].rect().bottomLeft()))
            diagnostics = tab.sections.widget(4).findChildren(QTableWidget)[-1]
            names = [diagnostics.item(i, 0).text() for i in range(diagnostics.rowCount())]
            assert "Просмотры с неоднозначным набором товаров" in names
            assert "Просмотры с ценой без определённой валюты" in names
    assert sample_result == original



def test_products_analytics_themes_headers_scroll_and_no_duplicates(window, sample_result):
    _, tab = window
    original = deepcopy(sample_result)
    group_headers = ["Товаров", "Просмотры", "Добавления в избранное", "Покупки", "Продано единиц"]
    expected = [
        ["Код", "Название", "Просмотры", "Клиенты с просмотрами", "Добавления в избранное", "Покупки"],
        ["Код", "Название", "Добавления в избранное", "Клиенты с избранным", "Просмотры", "Покупки"],
        ["Код", "Название", "Покупки", "Покупатели", "Продано единиц", "Просмотры", "Добавления в избранное"],
        list(ui.CATEGORY_HEADERS), *[["Значение", *group_headers] for _ in range(3)],
        ["Показатель", "Значение"],
        ["Система идентификаторов", "Взаимодействия", "Распознано", "Не распознано", "Доля распознанных, %"],
    ]
    for dark in (False, True):
        apply_app_theme(QApplication.instance(), dark)
        for _ in range(2):
            tab.render(sample_result)
            assert_section_names(tab)
            tab.sections.setCurrentIndex(2)
            QApplication.processEvents()
            page = tab.sections.widget(2)
            assert card_titles(page) == ["Товары с взаимодействиями", "Товары с просмотрами", "Товары в избранном", "Купленные товары"]
            assert [label.text() for label in page.findChildren(QLabel) if label.property("class") == "statisticsNumber"] == ["1", "1", "0", "1"]
            assert [label.text() for label in page.findChildren(QLabel) if label.property("class") == "statisticsSection"] == [
                "ПОПУЛЯРНЫЕ ТОВАРЫ ПО ПРОСМОТРАМ", "ПОПУЛЯРНЫЕ ТОВАРЫ ПО ИЗБРАННОМУ", "ПОПУЛЯРНЫЕ ТОВАРЫ ПО ПОКУПКАМ",
                "КАТЕГОРИИ ТОВАРОВ", "СПРОС ПО ПОЛУ ТОВАРА", "СПРОС ПО СЕЗОНУ", "СПРОС ПО СТИЛЕВОЙ ГРУППЕ", "КАЧЕСТВО СОПОСТАВЛЕНИЯ"]
            tables = page.findChildren(QTableWidget)
            assert [[table.horizontalHeaderItem(i).text() for i in range(table.columnCount())] for table in tables] == expected
            assert tables[2].item(0, 4).text() == "4"
            assert tables[3].item(0, 0).text() == "Не указан"
            assert tables[3].item(0, 1).text() == "Не указано"
            assert all(tables[i].item(0, 0).text() == "Не указано" for i in range(4, 7))
            assert tables[7].rowCount() == 5 and tables[7].item(2, 1).text() == "3"
            assert tables[8].item(0, 0).text() == "offline1C"
            assert page.widgetResizable() and page.widget().layout().spacing() == ui.SECTION_SPACING
            for table in tables:
                assert all(table.item(r, c).toolTip() == "" for r in range(table.rowCount()) for c in range(table.columnCount()))
            wait_until(lambda: page.verticalScrollBar().maximum() > 0)
            page.verticalScrollBar().setValue(page.verticalScrollBar().maximum())
            QApplication.processEvents()
            assert page.viewport().rect().contains(tables[-1].mapTo(page.viewport(), tables[-1].rect().bottomLeft()))
    assert original == sample_result


def test_ui_cleanup_preserves_diagnostics_and_styles_sections(window, sample_result):
    _, tab = window
    hidden = "Позиции заказов с нецелым количеством исключены из статистики."
    sample_result["warnings"] = (hidden, "Другое предупреждение.")
    sample_result["diagnostics"] = tuple((key, 7 if key == "fractional_quantity_order_lines" else value)
                                         for key, value in sample_result["diagnostics"])
    original = deepcopy(sample_result)
    for dark in (False, True):
        apply_app_theme(QApplication.instance(), dark)
        tab.render(sample_result)
        assert tab.status.text() == ui._calculation_message(sample_result) + "\nДругое предупреждение."
        for index in range(5):
            tab.sections.setCurrentIndex(index)
            QApplication.processEvents()
            page = tab.sections.widget(index)
            for label in page.findChildren(QLabel):
                assert not label.text().startswith(("Месяц определяется", "Показатели рассчитаны", "Показатели объединены",
                                                    "Бизнес-показатели", "Доступность и цена", "Активными считаются",
                                                    "Пол и возраст рассчитаны", "Если при чтении данных"))
                assert label.property("class") != "sectionHeader"
                if label.property("class") == "statisticsSection":
                    assert label.alignment() & Qt.AlignmentFlag.AlignHCenter
                    assert label.font().italic() and not label.font().bold()
                    assert label.font().pointSize() == 12
                    assert label.text() == label.text().upper()
            for table in page.findChildren(QTableWidget):
                assert table.height() <= 450
        diagnostics = tab.sections.widget(4).findChildren(QTableWidget)[-1]
        rows = {diagnostics.item(r, 0).text(): diagnostics.item(r, 1).text() for r in range(diagnostics.rowCount())}
        assert rows["Позиции с нецелым количеством"] == "7"
    assert sample_result == original


@pytest.mark.parametrize("count", [0, 1, 4, 60])
def test_table_content_height_and_scroll_after_resize_and_theme_change(app, count):
    window = QWidget()
    layout = QVBoxLayout(window)
    table = ui._table(layout, ["Колонка", "Значение"], [(f"Строка {i}", i) for i in range(count)])
    window.show()
    try:
        for dark in (False, True):
            apply_app_theme(app, dark)
            for width in (260, 1200, 260):
                window.resize(width, 800)
                table.setColumnWidth(0, 500)
                table.setColumnWidth(1, 200)
                QTest.qWait(30)
                if count < 60:
                    assert table.verticalScrollBar().maximum() == 0
                    assert abs(table.viewport().height() - sum(table.rowHeight(r) for r in range(count))) <= 2
                else:
                    assert table.height() == 450
                    assert table.verticalScrollBar().maximum() > 0
                    table.scrollToBottom()
                    QTest.qWait(20)
                    assert 0 <= table.visualItemRect(table.item(count - 1, 0)).bottom() < table.viewport().height()
                assert (table.horizontalScrollBar().maximum() > 0) == (width == 260)
    finally:
        window.close()
        window.deleteLater()
        app.processEvents()


def test_orders_cards_rates_full_stores_and_shared_scrollbars(window, sample_result):
    from PyQt6.QtCore import QEvent
    from PyQt6.QtGui import QColor
    from PyQt6.QtWidgets import QStyle, QStyleOptionSlider
    _, tab = window
    sample_result["orders_without_purchase"] = 2
    sample_result["purchase_basket_distribution"] = tuple((label, 1, rate) for label, rate in zip(
        ("1 позиция", "2 позиции", "3–5 позиций", "6–10 позиций", "11+ позиций"), (25., 15., 30., 20., 10.)))
    sample_result["store_statistics"] = tuple((currency, str(i), f"Магазин {i}", 1, 1, 1, "1", "1", "1")
                                               for currency in ("RUB", "KZT") for i in range(25))
    original = deepcopy(sample_result)
    for dark in (False, True):
        apply_app_theme(QApplication.instance(), dark)
        tab.render(sample_result)
        tab.sections.setCurrentIndex(1)
        QTest.qWait(30)
        page = tab.sections.widget(1)
        assert card_titles(page) == ["Заказы с покупкой", "Невыкупленные заказы", "Позиции покупок", "Покупатели",
                                    "Среднее число позиций на заказ", "Медиана числа позиций на заказ",
                                    "Доля заказов с 1 позицией", "Доля заказов с 3+ позициями"]
        assert [label.text() for label in page.findChildren(QLabel) if label.property("class") == "statisticsNumber"] == [
            "1", "2", "1", "1", "1.00", "1.00", "25.00%", "60.00%"]
        tables = page.findChildren(QTableWidget)
        for table in tables[3:5]:
            assert table.rowCount() == 25 and table.height() == 450
            assert table.verticalScrollBar().maximum() > 0
        # Check the actual unhovered pixels, including internal horizontal scrolling.
        table = tables[3]
        table.setColumnWidth(0, 1500)
        page.ensureWidgetVisible(table)
        QTest.qWait(30)
        for scrollbar in (page.verticalScrollBar(), table.verticalScrollBar(), table.horizontalScrollBar()):
            QApplication.sendEvent(scrollbar, QEvent(QEvent.Type.Leave))
            option = QStyleOptionSlider()
            # Internal Qt scrollbars do not expose protected initStyleOption.
            option.initFrom(scrollbar)
            option.orientation = scrollbar.orientation()
            option.minimum = scrollbar.minimum()
            option.maximum = scrollbar.maximum()
            option.sliderPosition = scrollbar.sliderPosition()
            option.sliderValue = scrollbar.value()
            option.singleStep = scrollbar.singleStep()
            option.pageStep = scrollbar.pageStep()
            option.upsideDown = scrollbar.invertedAppearance()
            option.subControls = QStyle.SubControl.SC_All
            option.activeSubControls = QStyle.SubControl.SC_None
            if option.orientation == Qt.Orientation.Horizontal:
                option.state |= QStyle.StateFlag.State_Horizontal
                option.upsideDown ^= scrollbar.layoutDirection() == Qt.LayoutDirection.RightToLeft
            else:
                option.state &= ~QStyle.StateFlag.State_Horizontal
            option.state &= ~QStyle.StateFlag.State_MouseOver
            rect = scrollbar.style().subControlRect(QStyle.ComplexControl.CC_ScrollBar, option,
                                                    QStyle.SubControl.SC_ScrollBarSlider, scrollbar)
            assert rect.isValid() and not rect.isEmpty()
            assert rect.intersects(scrollbar.rect())
            assert scrollbar.rect().contains(rect.center())
            pixmap = scrollbar.grab()
            image = pixmap.toImage()
            ratio = pixmap.devicePixelRatio()
            x, y = round(rect.center().x() * ratio), round(rect.center().y() * ratio)
            assert 0 <= x < image.width() and 0 <= y < image.height()
            assert image.pixelColor(x, y) == QColor(os.environ['QTMATERIAL_PRIMARYCOLOR'])
    assert sample_result == original
    sample_result["purchase_basket_distribution"] = tuple((label, 0, 0.) for label, _, _ in original["purchase_basket_distribution"])
    tab.render(sample_result)
    numbers = [label.text() for label in tab.sections.widget(1).findChildren(QLabel) if label.property("class") == "statisticsNumber"]
    assert numbers[-2:] == ["0.00%", "0.00%"]


def test_category_code_name_legacy_path_and_final_headers(window, sample_result):
    _, tab = window
    sample_result['product_category_statistics'] = (('0638', 'Рубашки', 1, 2, 0, 1, '4'),)
    for dark in (False, True):
        apply_app_theme(QApplication.instance(), dark)
        tab.render(sample_result)
        tab.sections.setCurrentIndex(2)
        QTest.qWait(30)
        page = tab.sections.widget(2)
        category = page.findChildren(QTableWidget)[3]
        assert category.columnCount() == 8
        assert [category.item(0, i).text() for i in range(8)] == [
            '0638', 'Рубашки', 'Путь недоступен в сохранённом результате', '1', '2', '0', '1', '4']
        headings = [label for label in page.findChildren(QLabel) if label.property('class') == 'statisticsSection']
        assert len(headings) == 8
        for label in headings:
            assert label.height() <= label.fontMetrics().lineSpacing() + 2
        assert card_titles(tab.sections.widget(3))[3] == 'Повторные покупатели (от 2 покупок)'
        orders = [label.text() for label in tab.sections.widget(1).findChildren(QLabel) if label.property('class') == 'statisticsSection']
        assert 'МАГАЗИНЫ И КАНАЛЫ — RUB' in orders and 'МАГАЗИНЫ И КАНАЛЫ — KZT' in orders
        assert not any('ПО РОССИИ' in value or 'ПО КАЗАХСТАНУ' in value for value in orders)


def test_export_cache_result_and_dialog_cancel(window, sample_result, monkeypatch):
    parent, tab = window
    cache.save_result(tab.cache_path, sample_result)
    restored = ui.DatasetStatisticsTab(parent)
    assert restored.displayed_result is not None and restored.export_button.isEnabled()
    restored.deleteLater()
    tab.render(sample_result)
    status, page_status = parent.status_label.text(), tab.status.text()
    calls = []
    def dialog(*args):
        assert args[1] == 'Экспорт статистики' and args[3] == 'ZIP (*.zip)'
        assert args[2].endswith('Статистика_2026-09-24_09-47-28.zip')
        return '', ''
    monkeypatch.setattr(ui.QFileDialog, 'getSaveFileName', dialog)
    monkeypatch.setattr(ui, 'export_statistics_bundle', lambda *args: calls.append(args))
    tab.export_button.click()
    assert not calls and tab._export_task is None
    assert parent.status_label.text() == status and tab.status.text() == page_status
    assert tab.export_button.isEnabled()


@pytest.mark.parametrize('suffix', ['', '.zip', '.ZIP'])
def test_export_uses_displayed_snapshot_despite_new_filter(window, sample_result, monkeypatch, tmp_path, suffix):
    from Application.analysis_filter import AnalysisFilter
    from Application.statistics_bundle_export import export_statistics_bundle
    from openpyxl import load_workbook
    parent, tab = window
    tab.render(sample_result)
    cache.save_result(tab.cache_path, sample_result)
    before = tab.cache_path.read_bytes()
    tab.set_analysis_filter(AnalysisFilter(nomenclature_types=('Новый фильтр',)))
    tab.pending = {**sample_result, 'total_interactions': 999}
    status = tab.status.text()
    destination = tmp_path / ('export' + suffix)
    monkeypatch.setattr(ui.QFileDialog, 'getSaveFileName', lambda *args: (str(destination), 'ZIP (*.zip)'))
    calls = []
    def save(result, path):
        assert result is tab.displayed_result is sample_result
        calls.append(path)
        export_statistics_bundle(result, path)
    monkeypatch.setattr(ui, 'export_statistics_bundle', save)
    tab.export_button.click()
    wait_until(lambda: tab._export_task is None)
    assert len(calls) == 1 and calls[0] == str(destination) + ('' if suffix else '.zip')
    assert parent.status_label.text() == 'Статистика экспортирована в ZIP: Excel + HTML'
    assert parent._status_reset_timer.interval() == 5000
    assert tab.status.text() == status and tab.cache_path.read_bytes() == before
    assert tab.export_button.isEnabled()
    from io import BytesIO
    from zipfile import ZipFile
    with ZipFile(calls[0]) as archive:
        assert archive.testzip() is None and len(archive.namelist()) == 2
        html = archive.read(next(name for name in archive.namelist() if name.endswith('.html'))).decode('utf-8')
        assert 'Новый фильтр' not in html
        book = load_workbook(BytesIO(archive.read(next(name for name in archive.namelist() if name.endswith('.xlsx')))))
    metadata = {row[0]: row[1] for row in book['Сводка'].values if row[0]}
    assert metadata['Вид номенклатуры'] == 'Все значения'
    # Ordinary KPI and monthly tables belong to the displayed snapshot too.
    from test_statistics_export import kpi, find, assert_minimal_tables
    from decimal import Decimal
    assert kpi(book['Сводка'], 'Количество взаимодействий').value == sample_result['total_interactions']
    assert assert_minimal_tables(book)
    for month, _, rub, _, kzt in sample_result['order_monthly_dynamics']:
        row = find(book['Заказы'], month[5:] + '.' + month[:4])
        assert Decimal(str(row[2].value)) == Decimal(rub)
        assert Decimal(str(row[4].value)) == Decimal(kzt)
    book.close()


@pytest.mark.parametrize('error', [PermissionError, RuntimeError])
def test_export_error_preserves_snapshot_cache_and_page_status(window, sample_result, monkeypatch, tmp_path, error):
    parent, tab = window
    tab.render(sample_result)
    cache.save_result(tab.cache_path, sample_result)
    previous = tab.cache_path.read_bytes(), tab.status.text(), deepcopy(sample_result)
    monkeypatch.setattr(ui.QFileDialog, 'getSaveFileName', lambda *args: (str(tmp_path / 'export.xlsx'), ''))
    def failed(*args):
        raise error('synthetic')
    monkeypatch.setattr(ui, 'export_statistics_bundle', failed)
    tab.export_button.click()
    wait_until(lambda: tab._export_task is None)
    assert parent.status_label.text() == 'Не удалось экспортировать статистику'
    assert parent._status_reset_timer.interval() == 5000
    assert tab.displayed_result is sample_result
    assert (tab.cache_path.read_bytes(), tab.status.text(), sample_result) == previous
    assert tab.export_button.isEnabled()


def test_export_contains_every_visible_table_and_card(window, sample_result, tmp_path):
    from Application.statistics_export import export_statistics, SHEET_NAMES
    from openpyxl import load_workbook
    _, tab = window
    tab.render(sample_result)
    path = tmp_path / 'complete.xlsx'
    export_statistics(sample_result, path)
    book = load_workbook(path)
    for index, name in enumerate(SHEET_NAMES[1:6]):
        sheet = book[name]
        rows = list(sheet.values)
        if name == 'Клиенты':
            rows.extend(book['RFM-анализ'].values)  # GUI subsections have their own XLSX sheets.
            rows.extend(book['Когортный анализ'].values)
        page = tab.sections.widget(index)
        for table in page.findChildren(QTableWidget):
            headers = tuple(table.horizontalHeaderItem(i).text() for i in range(table.columnCount()))
            assert any(row[:len(headers)] == headers for row in rows), (name, headers)
        for title in card_titles(page):
            assert any(row[0] == title for row in rows), (name, title)
    book.close()


def test_export_runs_in_background_and_holds_original_snapshot(window, sample_result, monkeypatch, tmp_path):
    from threading import Event, get_ident
    _, tab = window
    tab.render(sample_result)
    entered, release = Event(), Event()
    gui_thread = get_ident()
    calls = []
    def save(result, path):
        assert get_ident() != gui_thread
        calls.append(result)
        entered.set()
        assert release.wait(10)
    monkeypatch.setattr(ui, 'export_statistics_bundle', save)
    monkeypatch.setattr(ui.QFileDialog, 'getSaveFileName', lambda *args: (str(tmp_path / 'export.xlsx'), ''))
    try:
        tab.export_button.click()
        wait_until(entered.is_set)
        assert not tab.export_button.isEnabled()
        tab.export()
        tab.render(deepcopy(sample_result))
        assert calls == [sample_result] and calls[0] is not tab.displayed_result
    finally:
        release.set()
    wait_until(lambda: tab._export_task is None)
    assert tab.export_button.isEnabled()


def test_export_worker_can_finish_after_tab_destruction(window, sample_result, monkeypatch, tmp_path):
    from threading import Event
    from PyQt6 import sip
    from PyQt6.QtCore import QCoreApplication, QEvent, QThreadPool

    _, tab = window
    tab.render(sample_result)
    entered, release, finished = Event(), Event(), Event()

    def save(result, path):
        entered.set()
        assert release.wait(10)
        finished.set()
    monkeypatch.setattr(ui, 'export_statistics_bundle', save)
    monkeypatch.setattr(ui.QFileDialog, 'getSaveFileName', lambda *args: (str(tmp_path / 'export.zip'), ''))
    try:
        tab.export()
        wait_until(entered.is_set)
        tab.deleteLater()
        QCoreApplication.sendPostedEvents(None, QEvent.Type.DeferredDelete)
        assert sip.isdeleted(tab)
    finally:
        release.set()
    wait_until(finished.is_set)
    assert QThreadPool.globalInstance().waitForDone(5000)

@pytest.mark.parametrize('filtered', [False, True])
def test_system_action_types_stay_technical_and_total_card_is_unchanged(window, sample_result, filtered):
    from Application.analysis_filter import AnalysisFilter
    _, tab = window
    sample_result['source_actions'] = 100
    sample_result['actions_with_product'] = 2
    sample_result['actions_without_product'] = 98
    sample_result['action_types'] = (('ProsmotrProdukta', 90), ('ProsmotrKategoriiProduktov', 10))
    sample_result['diagnostics'] = tuple((key, 88 if key == 'mapped_without_product' else value)
                                         for key, value in sample_result['diagnostics'])
    if filtered:
        sample_result['analysis_filter'] = AnalysisFilter(nomenclature_types=('Рубашки',)).to_dict()
    tab.render(sample_result)
    assert tab.card_labels[0].text() == '3'
    actions = tab.sections.widget(0)
    assert not any(t.horizontalHeaderItem(0).text() in ('Системное название', 'Классификация событий')
                   for t in actions.findChildren(QTableWidget))
    technical = tab.sections.widget(4)
    table = next(t for t in technical.findChildren(QTableWidget) if t.horizontalHeaderItem(0).text() == 'Системное название')
    assert [float(table.item(i, 2).text()) for i in range(2)] == [90., 10.]
    quality = next(t for t in technical.findChildren(QTableWidget) if t.horizontalHeaderItem(0).text() == 'Классификация событий')
    rows = {quality.item(i, 0).text(): quality.item(i, 1).text() for i in range(quality.rowCount())}
    assert rows[ui.DIAGNOSTIC_LABELS['mapped_without_product']] == '88'
    assert tab.export_button.isEnabled()


@pytest.mark.parametrize('scenario', ['control', 'filtered', 'empty', 'categories_new', 'categories_legacy',
                                     'categories_unavailable', 'categories_empty', 'categories_hostile',
                                     'scope_new', 'scope_missing_profiles', 'scope_empty', 'scope_legacy', 'scope_filtered'])
def test_saved_quality_and_categories_gui_xlsx_html_parity_without_source_reads(
        window, sample_result, tmp_path, monkeypatch, scenario):
    import sqlite3
    from decimal import Decimal
    from openpyxl import load_workbook
    from Application import dataset_statistics, statistics_slices, statistics_html_export as html
    from Application.mindbox import raw_reader
    from Application.statistics_export import export_statistics, PERCENT_FORMAT
    from Application.statistics_presentation import (
        SOURCE_EVENT_QUALITY_HEADING, SOURCE_EVENT_QUALITY_NOTES, RESOLUTION_NOTES, source_event_quality_rows,
    )
    from statistics_export_samples import quality_snapshot, empty_snapshot, category_snapshot, customer_scope_snapshot
    from Application.analysis_filter import AnalysisFilter

    _, tab = window
    result = quality_snapshot(sample_result)
    if scenario.startswith('scope_'):
        result = customer_scope_snapshot(sample_result, missing_profiles=scenario == 'scope_missing_profiles',
                                         empty=scenario == 'scope_empty')
        if scenario == 'scope_legacy':
            result.pop('customer_scope')
        if scenario == 'scope_filtered':
            result['analysis_filter'] = AnalysisFilter(nomenclature_types=('Рубашки',)).to_dict()
    if scenario.startswith('categories_'):
        result = category_snapshot(sample_result)
        if scenario == 'categories_legacy':
            result.pop('product_category_paths')
        elif scenario == 'categories_unavailable':
            result['product_category_paths'] = tuple((code, None) for code, _ in result['product_category_paths'])
        elif scenario == 'categories_hostile':
            code = '=1+1 ё–<tag>'
            categories = result['product_category_statistics']
            result['product_category_statistics'] = ((code, *categories[0][1:]), categories[1])
            result['product_category_paths'] = ((code, ('<img src=x onerror=alert(1)>', 'Рубашки')),
                                                result['product_category_paths'][1])
    if scenario == 'filtered':
        result['analysis_filter'] = AnalysisFilter(nomenclature_types=('Рубашки',)).to_dict()
    elif scenario in ('empty', 'categories_empty'):
        result = empty_snapshot(result)
    path = tmp_path / 'saved.json'
    cache.save_result(path, result)
    saved_bytes = path.read_bytes()
    cached = cache.load_result(path)
    original = deepcopy(cached)

    def forbidden(*args, **kwargs):
        raise AssertionError('Presentation attempted a source read or recalculation')

    monkeypatch.setattr(dataset_statistics, 'calculate_statistics', forbidden)
    monkeypatch.setattr(dataset_statistics, 'iter_export', forbidden)
    monkeypatch.setattr(raw_reader, 'iter_export', forbidden)
    monkeypatch.setattr(statistics_slices, 'calculate_cached_statistics', forbidden)
    monkeypatch.setattr(sqlite3, 'connect', forbidden)
    original_open = Path.open

    def open_without_sources(path, *args, **kwargs):
        if path.resolve().is_relative_to(ui.PROJECT_ROOT / 'input_data'):
            forbidden()
        return original_open(path, *args, **kwargs)

    monkeypatch.setattr(Path, 'open', open_without_sources)
    tab.render(cached)
    tab.sections.setCurrentIndex(0)
    QApplication.processEvents()
    rows = source_event_quality_rows(cached)
    page = tab.sections.widget(0)
    assert any(label.text() == SOURCE_EVENT_QUALITY_HEADING.upper() for label in page.findChildren(QLabel))
    quality = next(table for table in page.findChildren(QTableWidget)
                   if table.horizontalHeaderItem(0).text() == 'Показатель')
    assert [(quality.item(i, 0).text(), quality.item(i, 1).text()) for i in range(quality.rowCount())] == [
        (label, 'Нет данных' if value is None else ui._number(value)) for label, value in rows[:6] + rows[7:]]
    rate = rows[6][1]
    expected_text = 'Нет данных' if rate is None else f'{rate:.2f}'.replace('.', ',') + ' %'
    assert [label.text() for label in page.findChildren(QLabel)
            if label.property('class') == 'statisticsNumber'][-1] == expected_text
    assert all(text in '\n'.join(label.text() for label in page.findChildren(QLabel))
               for _, text in SOURCE_EVENT_QUALITY_NOTES)
    products = tab.sections.widget(2)
    assert ' '.join(text for _, text in RESOLUTION_NOTES) in [label.text() for label in products.findChildren(QLabel)]
    category_rows = ui.product_category_rows(cached)
    category_table = products.findChildren(QTableWidget)[3]
    assert [category_table.horizontalHeaderItem(i).text() for i in range(8)] == list(ui.CATEGORY_HEADERS)
    assert category_table.rowCount() == len(category_rows)
    assert [tuple(category_table.item(r, c).text() for c in range(8)) for r in range(category_table.rowCount())] == [
        (*row[:3], *(ui._number(v) for v in row[3:7]), ui._decimal_display(row[7], money=False)) for row in category_rows]
    assert ui.DIRECT_CATEGORY_NOTE in [label.text() for label in products.findChildren(QLabel)]
    scope_sections = ui.customer_scope_sections(cached)
    client_tables = tab.sections.widget(3).findChildren(QTableWidget)
    for table, (_, values) in zip(client_tables[:3], scope_sections):
        assert [(table.item(r, 0).text(), table.item(r, 1).text()) for r in range(table.rowCount())] == [
            (label, ui._number(value)) for label, value in values]
        assert all(table.rowHeight(row) >= table.sizeHintForRow(row) for row in range(table.rowCount()))
    client_text = '\n'.join(label.text() for label in tab.sections.widget(3).findChildren(QLabel))
    assert ui.customer_scope_context(cached) in client_text
    assert all(text in client_text for _, text in ui.CUSTOMER_SCOPE_NOTES)
    if scenario == 'categories_new':
        assert category_table.columnWidth(2) >= category_table.fontMetrics().horizontalAdvance(category_rows[0][2])
        assert category_rows[0][0] == '000123' and category_rows[1][0] == '000124'
        assert category_rows[0][1] == category_rows[1][1] == 'Рубашки'

    xlsx = tmp_path / 'quality.xlsx'
    export_statistics(cached, xlsx)
    book = load_workbook(xlsx)
    try:
        sheet_rows = list(book['Действия'].values)
        for label, value in rows:
            row = next(row for row in sheet_rows if row[0] == label)
            if label == rows[6][0] and value is not None:
                assert row[1] == pytest.approx(value)
            else:
                assert row[1] == ('Нет данных' if value is None else value)
        rate_cell = next(row[1] for row in book['Действия'].iter_rows() if row[0].value == rows[6][0])
        assert rate_cell.data_type == ('s' if rate is None else 'n')
        assert rate_cell.number_format == ('@' if rate is None else PERCENT_FORMAT)
        assert all(any(row[:2] == pair for row in sheet_rows) for pair in SOURCE_EVENT_QUALITY_NOTES)
        assert all(any(row[:2] == pair for row in book['Товары'].values) for pair in RESOLUTION_NOTES)
        for row in category_rows:
            cells = next(cells for cells in book['Товары'].iter_rows()
                         if tuple(cell.value for cell in cells[:3]) == row[:3])
            assert cells[0].data_type == 's' and cells[0].number_format == '@'
            assert tuple(cell.value for cell in cells[:7]) == row[:7]
            assert Decimal(str(cells[7].value)) == Decimal(row[7])
        client_rows = list(book['Клиенты'].values)
        for _, values in scope_sections:
            for label, value in values:
                row = next(row for row in client_rows if row[0] == label)
                assert row[1] == value
    finally:
        book.close()

    sections = html.presentation(cached)
    scope_html_rows = [row for block in sections[4].blocks if block['kind'] == 'table' for row in block['rows']]
    for _, values in scope_sections:
        for label, value in values:
            cells = next(row for row in scope_html_rows if row[0]['text'] == html.display_text(label))
            if isinstance(value, (int, float)):
                assert Decimal(cells[1]['sort']) == Decimal(str(value))
            else:
                assert cells[1]['text'] == html.display_text(value)
    category_block = next(block for block in sections[3].blocks if block.get('headers') == list(ui.CATEGORY_HEADERS))
    for cells, row in zip(category_block['rows'], category_rows):
        assert cells[0]['text'] == row[0]  # Codes bypass display typography translation.
        assert cells[1]['text'] == html.display_text(row[1])
        assert cells[2]['text'] == html.display_text(row[2])
        assert tuple(Decimal(cell['sort']) for cell in cells[3:]) == tuple(Decimal(str(v)) for v in row[3:])
    html_rows = [row for block in sections[1].blocks if block['kind'] == 'table' for row in block['rows']]
    for label, value in rows:
        cells = next(row for row in html_rows if row[0]['text'] == html.display_text(label))
        assert cells[1]['text'] == ('Нет данных' if value is None else html.exact_number(value)
                                   + (' %' if label == rows[6][0] else ''))
        if value is not None:
            assert Decimal(cells[1]['sort']) == Decimal(str(value))
    document = html.build_statistics_html(cached)
    assert html.display_text(SOURCE_EVENT_QUALITY_HEADING) in document
    assert all(html.display_text(text) in document for _, text in SOURCE_EVENT_QUALITY_NOTES + RESOLUTION_NOTES)
    assert html.display_text(ui.DIRECT_CATEGORY_NOTE) in document
    if scenario == 'categories_hostile':
        assert '<img src=x onerror=alert(1)>' not in document
        assert '&lt;img src=x onerror=alert(1)&gt;' in document
        assert html._cell(code, identifier=True)['text'] == code
    assert tab.displayed_result is cached
    assert cached == original
    assert path.read_bytes() == saved_bytes
    assert cached['resolution_rate'] == (0. if scenario in ('empty', 'categories_empty', 'scope_empty', 'scope_missing_profiles') else 100.)


def test_final_chart_set_preserves_tables_and_section_titles(window, sample_result):
    from Application.statistics_charts import StatisticsChart
    from test_statistics_charts import _populated
    from PyQt6.QtCore import QCoreApplication, QEvent
    _, tab = window
    result = _populated(sample_result)
    for _ in range(3):
        tab.render(result)
        QCoreApplication.sendPostedEvents(None, QEvent.Type.DeferredDelete)
        QApplication.processEvents()
        assert [len(tab.sections.widget(i).findChildren(StatisticsChart)) for i in range(5)] == [4, 6, 4, 3, 0]
        assert [len(tab.sections.widget(i).findChildren(QTableWidget)) for i in range(5)] == [8, 7, 9, 12, 5]
        chart_widgets = tab.findChildren(StatisticsChart)
        assert len(chart_widgets) == 17
        actions = _visual_items(tab.sections.widget(0).widget().layout())
        action_charts = tab.sections.widget(0).findChildren(StatisticsChart)
        assert {chart.data.title for chart in action_charts} == {
            'Распределение клиентов по количеству просмотров', 'Распределение клиентов по количеству избранного',
            'Динамика просмотров', 'Динамика избранного'}
        for kind, header in [('просмотров', 'Количество просмотров'), ('избранного', 'Количество избранного')]:
            chart = next(chart for chart in action_charts if chart.data.title.endswith(kind))
            index = next(i for i, item in enumerate(actions) if item.widget() is chart)
            assert actions[index - 1].widget().text() == chart.data.title.upper()
            following = actions[index + 1].widget()
            assert isinstance(following, QTableWidget) and following.horizontalHeaderItem(0).text() == header
        view = next(chart for chart in action_charts if chart.data.title.endswith('просмотров'))
        favorite = next(chart for chart in action_charts if chart.data.title.endswith('избранного'))
        assert next(i for i, item in enumerate(actions) if item.widget() is favorite) == next(i for i, item in enumerate(actions) if item.widget() is view) + 3
        assert not any(label.text() == 'ДИНАМИКА КЛИЕНТОВ С ПРОСМОТРАМИ' for label in tab.findChildren(QLabel))
        assert sum(len(chart.findChildren(QLabel)) for chart in chart_widgets) == 17
        clients = _visual_items(tab.sections.widget(3).widget().layout())
        frequency = next(chart for chart in chart_widgets if chart.data.title == 'Распределение клиентов по количеству взаимодействий')
        index = next(i for i, item in enumerate(clients) if item.widget() is frequency)
        following = clients[index + 1].widget()
        assert isinstance(following, QTableWidget)
        assert following.horizontalHeaderItem(0).text() == 'Количество взаимодействий'
        assert sorted(chart.chart().title() for chart in chart_widgets if chart.chart().title()) == [
            'ДИНАМИКА ИЗБРАННОГО', 'ДИНАМИКА ПРОСМОТРОВ']
        for chart in chart_widgets:
            page = next(tab.sections.widget(i) for i in range(4) if chart in tab.sections.widget(i).findChildren(StatisticsChart))
            assert page.horizontalScrollBar().maximum() == 0
