"""RFM display/export reads only saved aggregate snapshots, in both themes."""

from copy import deepcopy
from decimal import Decimal
from pathlib import Path
import sqlite3

import pytest
from openpyxl import load_workbook
from PyQt6.QtWidgets import QApplication, QLabel, QTableWidget

from Application import dataset_statistics as statistics, statistics_cache as cache, statistics_slices as slices
from Application import statistics_export as xlsx, statistics_html_export as html
from Application.statistics_rfm_core import RfmAggregates
from Application.statistics_rfm_presentation import rfm_sections, RFM_LEGACY_TEXT
from Application.statistics_presentation import buyer_sections
from Application.tabs import dataset_statistics_tab as ui
from Application.theme.apply_theme import apply_app_theme
from tests.test_statistics_rfm_integration import dataset, calculate, SELECTION, save
from Application.analysis_filter import AnalysisFilter
from tests.test_statistics_rfm_core import fact as ready_fact, SPACE
from tests.tabs.test_dataset_statistics_tab import window, app


__all__ = ['dataset','window','app']
pytestmark = pytest.mark.usefixtures('window_settings')


@pytest.mark.parametrize('scenario',['full','partial','empty','no_orders','legacy'])
@pytest.mark.parametrize('dark',[False,True])
def test_rfm_saved_snapshot_gui_xlsx_html_identical_values_no_raw_access(dataset,window,tmp_path,monkeypatch,scenario,dark):
    result = calculate(dataset)
    if scenario=='legacy':
        result.pop('rfm_statistics')
    elif scenario=='no_orders':
        dataset[2]['manual_interactions'] = None
        dataset[2]['actions'] = dataset[2]['orders'] = {}
        save(dataset[0],dataset[2])
        result = calculate(dataset,AnalysisFilter())
    elif scenario in ('empty','partial'):
        coverage = [(cache.parse_timestamp(a),cache.parse_timestamp(b)) for a,b in
                    next(s['intervals'] for s in result['coverage'] if s['source']=='Orders')]
        if scenario=='partial':
            coverage = [(cache.parse_timestamp('2026-01-01T00:00:00Z'),cache.parse_timestamp('2026-09-01T00:00:00Z'))]
            next(s for s in result['coverage'] if s['source']=='Orders')['intervals'] = [(a.isoformat(),b.isoformat()) for a,b in coverage]
            result['buyer_statistics'] = None
        result['cohort_statistics'] = None  # This fixture replaces Orders coverage for RFM only.
        rfm = RfmAggregates(cache.parse_timestamp('2026-08-31T00:00:00Z').date(),coverage,identity_space=SPACE)
        if scenario=='partial':
            rfm.add(ready_fact())
        result['rfm_statistics'] = rfm.result()
    path = tmp_path/'rfm.json'
    cache.save_result(path,result)
    result = cache.load_result(path)
    before = deepcopy(result)
    saved = path.read_bytes()
    _,tab = window
    def forbidden(*args,**kwargs):
        raise AssertionError('RFM presentation attempted a source read or recalculation')
    with monkeypatch.context() as patch:
        patch.setattr(statistics,'iter_export',forbidden)
        patch.setattr(statistics,'calculate_statistics',forbidden)
        patch.setattr(slices,'calculate_cached_statistics',forbidden)
        patch.setattr(RfmAggregates,'result',forbidden)
        patch.setattr(sqlite3,'connect',forbidden)
        original_open = Path.open
        def guarded_open(p,*args,**kwargs):
            if p.resolve().is_relative_to(ui.PROJECT_ROOT/'input_data'):
                forbidden()
            return original_open(p,*args,**kwargs)
        patch.setattr(Path,'open',guarded_open)
        apply_app_theme(QApplication.instance(),dark)
        tab.render(result)
        tab.sections.setCurrentIndex(3)
        QApplication.processEvents()
        target = tmp_path/'rfm.xlsx'
        xlsx.export_statistics(result,target)
        rendered = html.build_statistics_html(result)
    assert result==before and path.read_bytes()==saved
    sections = rfm_sections(result)
    page = tab.sections.widget(3)
    labels = [label.text() for label in page.findChildren(QLabel)]
    assert all(title.upper() in labels for title,_,_,_ in sections)
    tables = page.findChildren(QTableWidget)
    start = 9+sum(len(t) for _,_,t,_ in buyer_sections(result))
    expected = [(headers,rows) for _,_,blocks,_ in sections for headers,rows in blocks]
    rfm_tables = tables[start:start+len(expected)]
    assert len(rfm_tables)==len(expected)
    for table,(headers,rows) in zip(rfm_tables,expected):
        assert [table.horizontalHeaderItem(c).text() for c in range(table.columnCount())]==list(headers)
        assert [[table.item(r,c).text() for c in range(table.columnCount())] for r in range(table.rowCount())]==[
            [ui._number(ui._buyer_display(v)) for v in row] for row in rows]
        assert all(table.rowHeight(r)>=table.sizeHintForRow(r) for r in range(table.rowCount()))
    book = load_workbook(target)
    try:
        assert 'RFM-анализ' in book.sheetnames
        rows = list(book['RFM-анализ'].values)
        html_rows = [row for block in html.presentation(result)[4].blocks if block['kind']=='table' for row in block['rows']]
        for title,cards,blocks,notes in sections:
            assert any(row[0]==title for row in rows)
            assert html.display_text(title) in rendered
            for label,value in cards:
                assert any(row[0]==label for row in rows)
            for headers,values in blocks:
                assert any(tuple(row[:len(headers)])==tuple(headers) for row in rows)
                for row in values:
                    shown = [html._cell(v)['text'] for v in row]
                    assert any([cell['text'] for cell in candidate]==shown for candidate in html_rows)
                    assert any(all(Decimal(str(actual))==pytest.approx(v) if isinstance(v,Decimal) else
                                   actual==pytest.approx(v) if isinstance(v,float) else actual==v
                                   for actual,v in zip(candidate,row)) for candidate in rows)
            assert all(html._cell(text)['text'] in rendered for _,text in notes)
        if scenario=='legacy':
            assert RFM_LEGACY_TEXT in rendered
        else:
            assert all(x not in rendered for x in ('"customer_id"','"order_id"','"identity_space"'))
    finally:
        book.close()


def test_old_cache_opens_on_startup_and_exports_without_recalculation(dataset,window,tmp_path,monkeypatch):
    result = calculate(dataset)
    result.pop('rfm_statistics')
    result.pop('cohort_statistics')
    path = tmp_path/'old-cache.json'
    cache.save_result(path,result)
    before = path.read_bytes()
    monkeypatch.setattr(ui,'load_filter',lambda:SELECTION)
    monkeypatch.setattr(ui,'CACHE_PATH',path)
    monkeypatch.setattr(ui.DatasetStatisticsTab,'start',lambda self:pytest.fail('Automatic recalculation'))
    tab = ui.DatasetStatisticsTab(window[0])
    assert tab.displayed_result is not None
    assert path.read_bytes()==before
    assert any(label.text()==RFM_LEGACY_TEXT for label in tab.findChildren(QLabel)) is False  # Table placeholder, no fake card.
    tables = tab.sections.widget(3).findChildren(QTableWidget)
    assert tables[-1].item(0,1).text()==RFM_LEGACY_TEXT
    tab._shutdown()
