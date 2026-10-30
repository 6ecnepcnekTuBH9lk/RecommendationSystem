"""Shared buyer GUI/XLSX/HTML values over saved synthetic snapshots only."""

from copy import deepcopy
from dataclasses import asdict
from decimal import Decimal
from pathlib import Path
import sqlite3
from Application.analysis_filter import AnalysisFilter

import pytest
from openpyxl import load_workbook
from PyQt6.QtWidgets import QApplication, QLabel, QTableWidget

from Application import dataset_statistics as statistics, statistics_cache as cache
from Application import statistics_export as xlsx, statistics_html_export as html
from Application import statistics_slices
from Application.mindbox import raw_reader
from Application.statistics_presentation import buyer_sections, BUYER_LEGACY_TEXT
from Application.tabs import dataset_statistics_tab as ui
from Application.theme.apply_theme import apply_app_theme
from tests.test_statistics_buyer_aggregates import dataset, AT, aggregate, purchase, save
from tests.tabs.test_dataset_statistics_tab import window, app


__all__ = ['dataset','window','app']
pytestmark = pytest.mark.usefixtures('window_settings')


@pytest.mark.parametrize('scenario',['new','empty','missing_profiles','partial_money','legacy','no_orders','no_coverage'])
@pytest.mark.parametrize('dark',[False,True])
def test_saved_buyer_blocks_gui_xlsx_html_parity_no_sources(dataset,window,tmp_path,monkeypatch,scenario,dark):
    base = asdict(statistics.calculate_statistics(raw_root=dataset[0],catalog_path=dataset[1],calculated_at=AT))
    if scenario == 'no_orders':
        dataset[2]['orders'] = {}
        save(dataset[0],dataset[2])
        base = asdict(statistics.calculate_statistics(raw_root=dataset[0],catalog_path=dataset[1],calculated_at=AT))
    elif scenario == 'no_coverage':
        base = asdict(statistics.calculate_statistics(raw_root=dataset[0],catalog_path=dataset[1],calculated_at=AT,
                                                       analysis_filter=AnalysisFilter('2024-01-01','2024-12-31')))
    elif scenario == 'legacy':
        base.pop('buyer_statistics')
    elif scenario != 'new':
        coverage = [(cache.parse_timestamp(a),cache.parse_timestamp(b)) for a,b in base['buyer_statistics']['coverage']]
        records = [] if scenario == 'empty' else [purchase('a','known'),
            purchase('a','unknown',lines=(('F',1,'20','unsupported'),)),
            purchase('b','zero',lines=(('F',1,'0','offline1C'),))] if scenario == 'partial_money' else [purchase()]
        agg = aggregate(records,coverage=coverage,profiles=scenario != 'missing_profiles')
        base['buyer_statistics'] = agg.result(AT.date())
        base['customer_scope'] = None
        if scenario == 'missing_profiles':
            base['customers'] = None
            # Existing legacy demographics use zero placeholders if source missing.
            for key in ('gender_distribution','age_distribution'):
                base[key] = [(label,0,0.) for label,_,_ in base[key]]
            base['mean_age'] = base['median_age'] = None
    path = tmp_path/'buyer-saved.json'
    cache.save_result(path,base)
    payload = path.read_bytes()
    result = cache.load_result(path)
    original = deepcopy(result)
    def forbidden(*args,**kwargs):
        raise AssertionError('Display/export attempted source access or calculation')
    with monkeypatch.context() as patch:
        patch.setattr(statistics,'calculate_statistics',forbidden)
        patch.setattr(statistics,'iter_export',forbidden)
        patch.setattr(raw_reader,'iter_export',forbidden)
        patch.setattr(statistics_slices,'calculate_cached_statistics',forbidden)
        patch.setattr(sqlite3,'connect',forbidden)
        old_open = Path.open
        def no_sources(p,*args,**kwargs):
            if p.resolve().is_relative_to(ui.PROJECT_ROOT/'input_data'):
                forbidden()
            return old_open(p,*args,**kwargs)
        patch.setattr(Path,'open',no_sources)
        _,tab = window
        apply_app_theme(QApplication.instance(),dark)
        tab.render(result)
        tab.sections.setCurrentIndex(3)
        QApplication.processEvents()
        xlsx_path = tmp_path/'buyers.xlsx'
        xlsx.export_statistics(result,xlsx_path)
        rendered = html.build_statistics_html(result)
    assert result == original and path.read_bytes() == payload
    sections = buyer_sections(result)
    page = tab.sections.widget(3)
    labels = [label.text() for label in page.findChildren(QLabel)]
    assert all(title.upper() in labels for title,_,_,_ in sections)
    buyer_table_count = sum(len(tables) for _,_,tables,_ in buyer_sections(result))
    gui_tables = page.findChildren(QTableWidget)[9:9+buyer_table_count]
    expected_tables = [(headers,rows) for _,_,tables,_ in sections for headers,rows in tables]
    assert len(gui_tables) == len(expected_tables)
    for table,(headers,rows) in zip(gui_tables,expected_tables):
        assert [table.horizontalHeaderItem(i).text() for i in range(table.columnCount())] == list(headers)
        assert [[table.item(r,c).text() for c in range(table.columnCount())] for r in range(table.rowCount())] == [
            [ui._number(ui._buyer_display(v)) for v in row] for row in rows]
        assert all(table.rowHeight(r) >= table.sizeHintForRow(r) for r in range(table.rowCount()))
    book = load_workbook(xlsx_path)
    try:
        sheet = list(book['Клиенты'].values)
        html_sections = html.presentation(result)[4]
        html_rows = [r for block in html_sections.blocks if block['kind']=='table' for r in block['rows']]
        for title,cards,tables,notes in sections:
            assert any(row[0] == title for row in sheet)
            assert html.display_text(title) in rendered
            for label,value in cards:
                row = next(row for row in sheet if row[0] == label)
                if isinstance(value,Decimal):
                    assert Decimal(str(row[1])) == pytest.approx(value)
                else:
                    assert row[1] == value
                assert any(cells[0]['text']==html.display_text(label) for cells in html_rows)
            for headers,rows in tables:
                assert any(tuple(row[:len(headers)])==tuple(headers) for row in sheet)
                for values in rows:
                    expected = [html._cell(v)['text'] for v in values]
                    assert any([cell['text'] for cell in candidate]==expected for candidate in html_rows)
            assert all(html._cell(text)['text'] in rendered for _,text in notes)
        if scenario == 'legacy':
            assert BUYER_LEGACY_TEXT in rendered
        elif scenario not in ('no_orders','no_coverage'):
            assert 'buyer-kpis' in rendered
            assert not any(token in rendered for token in ('"customer_id"','"order_id"','identity_space'))
    finally:
        book.close()
