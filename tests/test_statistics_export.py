"""Offline XLSX snapshot export, numeric types and atomic publication."""
from copy import deepcopy
from datetime import datetime
from decimal import Decimal
from pathlib import Path
from zipfile import ZipFile

from openpyxl import load_workbook
from openpyxl.utils.cell import range_boundaries
import pytest

from Application import statistics_export as export
from Application.analysis_filter import AnalysisFilter
from tabs import test_dataset_statistics_tab as samples
from statistics_export_samples import complete_snapshot, empty_snapshot, stress_snapshot

sample_result = samples.sample_result


def rows(sheet):
    return list(sheet.values)


def find(sheet, label):
    return next(row for row in sheet.iter_rows() if row[0].value == label)


def kpi(sheet, label):
    return find(sheet, label)[1]


def test_warnings_and_method_notes_are_not_exported_but_snapshot_is_preserved(sample_result, tmp_path):
    from Application.statistics_presentation import WARNING_LABELS
    warning = 'Есть конфликтующие снимки заказов: item interactions рассчитаны по первому снимку.'
    sample_result['warnings'] = (warning,)
    original = deepcopy(sample_result)
    path = tmp_path / 'conflict.xlsx'
    export.export_statistics(sample_result, path)
    book = load_workbook(path)
    try:
        text = '\n'.join(str(value) for sheet in book for row in sheet.values for value in row if value is not None)
        assert WARNING_LABELS[warning] not in text and 'Предупреждения' not in text
        assert 'Денежные значения хранятся числом.' not in text
        assert 'Данные профилей клиентов отсутствуют.' not in text
        assert sample_result == original
    finally:
        book.close()


def test_category_path_beyond_excel_limit_fails_atomically_without_silent_truncation(sample_result, tmp_path):
    from statistics_export_samples import category_snapshot
    from Application.statistics_html_export import presentation
    result = category_snapshot(sample_result)
    result['product_category_paths'] = (('000123', ('R' * 32768, 'Рубашки')), result['product_category_paths'][1])
    path = tmp_path / 'previous.xlsx'
    path.write_bytes(b'previous report')
    with pytest.raises(ValueError, match='32767'):
        export.export_statistics(result, path)
    assert path.read_bytes() == b'previous report'
    assert any('R' * 32768 in cell['text'] for block in presentation(result)[3].blocks
               if block['kind'] == 'table' for row in block['rows'] for cell in row)


def test_workbook_snapshot_metadata_types_and_sections(sample_result, tmp_path):
    result = sample_result
    result['analysis_filter'] = AnalysisFilter('2025-01-02', '2025-12-30', ('Рубашки', None), None).to_dict()
    result['warnings'] = ('Проверка предупреждения',)
    result['product_season_statistics'] = (('Зима 2025', 1, 2, 0, 1, '4'),)
    result['order_financials'] = (('RUB', 1, 1, '4', '3901955713.80', '3901955713.80', '3901955713.80'),
                                  ('KZT', 1, 1, '1', '123456.78', '123456.78', '123456.78'))
    result['store_statistics'] = tuple(('RUB', str(i), f'Магазин {i}', 1, 1, 1, '1', '10.20', '10.20') for i in range(25)) + (
        ('KZT', 'kz', 'Казахстан', 1, 1, 1, '1', '123456.78', '123456.78'),)
    original = deepcopy(result)
    path = tmp_path / 'statistics.xlsx'
    export.export_statistics(result, path)
    assert result == original
    book = load_workbook(path)
    assert_minimal_tables(book)
    summary = book['Сводка']
    assert find(summary, 'Дата и время расчета')[1].value == datetime.fromisoformat(result['calculated_at']).astimezone().strftime('%d.%m.%Y %H:%M:%S')
    assert find(summary, 'Вид номенклатуры')[1].value == 'Рубашки; Не указано'
    assert find(summary, 'Сезон')[1].value == 'Все значения'
    assert find(book['Техническая информация'], 'Дата начала')[1].value == datetime(2025, 1, 2)
    assert kpi(summary, 'Количество взаимодействий').value == 3
    assert all(cell.value != 'Проверка предупреждения' for row in summary for cell in row)
    orders = book['Заказы']
    for currency, expected in [('RUB', '3901955713.80'), ('KZT', '123456.78')]:
        row = find(orders, currency)
        assert row[4].data_type == 'n'
        assert Decimal(str(row[4].value)) == Decimal(expected)
        assert row[4].number_format == '#,##0.00'
    assert len([r for r in rows(orders) if isinstance(r[0], str) and r[0].startswith('Магазин ') and r[1] == 1]) == 25
    assert find(orders, 'Магазины и каналы — RUB') and find(orders, 'Магазины и каналы — KZT')
    assert find(book['Товары'], 'Зима 2025')
    assert find(book['Товары'], '000001')[0].data_type == 's'
    assert find(book['Клиенты'], 'Средний возраст')[1].value is None
    assert find(book['Клиенты'], 'Доля активных клиентов, совершивших покупку, %')[1].value == 50
    assert find(book['Техническая информация'], 'Источник')[2].value == 'Источник данных'
    assert find(book['Техническая информация'], 'Заказы вне периода статистики')[1].data_type == 'n'
    assert all(cell.value != 'Проверка предупреждения' for row in book['Техническая информация'] for cell in row)
    book.close()


def test_literal_labels_not_formulas_and_decimal_quantity(sample_result, tmp_path):
    sample_result['top_purchased_products'] = (('=1+1', 'Название', 1, 1, '1234.125', 2, 0),)
    path = tmp_path / 'literal.xlsx'
    export.export_statistics(sample_result, path)
    book = load_workbook(path)
    row = find(book['Товары'], '=1+1')
    assert row[0].data_type == 's'
    assert row[4].data_type == 'n' and Decimal(str(row[4].value)) == Decimal('1234.125')
    book.close()


@pytest.mark.parametrize('phase', ['table', 'save', 'fsync', 'replace'])
def test_atomic_failure_keeps_destination_and_cleans_temp(sample_result, tmp_path, monkeypatch, phase):
    path = tmp_path / 'existing.xlsx'
    path.write_bytes(b'previous file')
    def fail(*args):
        raise PermissionError('synthetic')
    if phase == 'table':
        monkeypatch.setattr(export._Sheet, 'table', fail)
    elif phase == 'save':
        monkeypatch.setattr(export.Workbook, 'save', fail)
    elif phase == 'fsync':
        monkeypatch.setattr(export.os, 'fsync', fail)
    else:
        monkeypatch.setattr(export.os, 'replace', fail)
    with pytest.raises(PermissionError):
        export.export_statistics(sample_result, path)
    assert path.read_bytes() == b'previous file'
    assert list(tmp_path.iterdir()) == [path]


def test_no_raw_or_qt_dependency(tmp_path):
    import subprocess
    import sys
    result = subprocess.run([sys.executable, '-c',
        "import sys; import Application.statistics_export; "
        "assert 'PyQt6' not in sys.modules; "
        "assert 'Application.dataset_statistics' not in sys.modules; "
        "assert 'Application.mindbox.raw_reader' not in sys.modules"], capture_output=True, text=True)
    assert result.returncode == 0, result.stderr


def test_raw_actions_in_technical_and_filter_independent(sample_result, tmp_path):
    sample_result['source_actions'] = 100
    sample_result['action_types'] = (('ProsmotrProdukta', 90), ('ProsmotrKategoriiProduktov', 10))
    raw_rows = []
    for filtered in (False, True):
        sample_result['analysis_filter'] = AnalysisFilter(nomenclature_types=('Рубашки',) if filtered else None).to_dict()
        path = tmp_path / f'raw-{filtered}.xlsx'
        export.export_statistics(sample_result, path)
        book = load_workbook(path)
        assert kpi(book['Сводка'], 'Количество взаимодействий').value == sample_result['total_interactions']
        business = {cell.value for row in book['Действия'] for cell in row}
        assert 'Системное название' not in business and 'Классификация событий' not in business
        technical = book['Техническая информация']
        assert find(technical, 'Исходные Actions')[1].value == 100
        raw = [find(technical, name) for name, _ in sample_result['action_types']]
        assert all(cell.data_type == 'n' for row in raw for cell in row[1:3])
        values = [(row[1].value, row[2].value) for row in raw]
        assert sum(v[1] for v in values) == pytest.approx(100)
        assert all(v[1] <= 100 for v in values)
        raw_rows.append(values)
        book.close()
    assert raw_rows[0] == raw_rows[1]


def assert_minimal_tables(book):
    """Saved workbook has only ordinary, bordered tables and no navigation."""
    assert book.sheetnames == list(export.SHEET_NAMES)
    names = set()
    checked = 0
    for sheet in book:
        assert sheet.sheet_state == 'visible'
        assert not sheet._charts and not sheet._images
        assert sheet.sheet_view.showGridLines is True
        assert sheet.freeze_panes is None
        assert sheet.sheet_view.pane is None
        assert not sheet.print_options.gridLines
        assert not sheet.protection.sheet
        assert not sheet.row_breaks.brk
        assert sheet.page_setup.fitToHeight == 0
        assert all(not column.hidden for column in sheet.column_dimensions.values())
        occupied = []
        for table in sheet.tables.values():
            assert table.name not in names
            names.add(table.name)
            left, top, right, bottom = range_boundaries(table.ref)
            assert all(bottom < a or b < top for a, b in occupied)
            occupied.append((top, bottom))
            assert not table.tableStyleInfo.showFirstColumn
            assert not table.tableStyleInfo.showColumnStripes
            assert not table.tableStyleInfo.showRowStripes
            for row in sheet.iter_rows(min_row=top, max_row=bottom, min_col=left, max_col=right):
                for cell in row:
                    for side in (cell.border.left, cell.border.right, cell.border.top, cell.border.bottom):
                        assert side.style == 'thin' and side.color.rgb[-6:] == export.BORDER_COLOR
                    if cell.row == top:
                        assert cell.fill.fgColor.rgb[-6:] == export.HEADER_FILL and cell.font.bold
                    else:
                        assert cell.fill.patternType is None
                    assert all(not (area.min_row <= cell.row <= area.max_row and
                                    area.min_col <= cell.column <= area.max_col)
                               for area in sheet.merged_cells.ranges)
            checked += 1
        assert all(cell.hyperlink is None and cell.data_type != 'f' for row in sheet for cell in row)
        assert all(cell.value not in ('Обзор', 'Таблицы', 'Аналитический обзор', 'Подпись диаграммы')
                   for row in sheet for cell in row)
        assert all((row.height or export.ROW_HEIGHT) <= export.MAX_ROW_HEIGHT
                   for row in sheet.row_dimensions.values())
        for row in sheet:
            for cell in row:
                if cell.value is not None:
                    assert cell.font.color.rgb[-6:] == export.INK
                    if cell.font.sz == 11 and cell.font.bold:
                        assert cell.fill.fgColor.rgb[-6:] == export.SECTION_FILL
    return checked


def test_saved_sheet_xml_has_grid_without_panes_or_off_table_styling(sample_result, tmp_path):
    from xml.etree import ElementTree as ET

    path = tmp_path / 'grid.xlsx'
    export.export_statistics(sample_result, path)
    namespace = {'m': 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
    book = load_workbook(path)
    try:
        assert assert_minimal_tables(book)
        with ZipFile(path) as archive:
            for index, sheet in enumerate(book, 1):
                xml = ET.fromstring(archive.read(f'xl/worksheets/sheet{index}.xml'))
                assert not xml.findall('.//m:pane', namespace)
                for view in xml.findall('.//m:sheetView', namespace):
                    assert view.get('showGridLines') == '1'
                for selection in xml.findall('.//m:selection', namespace):
                    assert selection.get('pane') in (None, 'topLeft')
                options = xml.find('m:printOptions', namespace)
                assert options is None or options.get('gridLines') in (None, '0')
                for dimension in xml.findall('.//m:col', namespace):
                    assert dimension.get('style') in (None, '0')
                for dimension in xml.findall('.//m:row', namespace):
                    assert dimension.get('s') in (None, '0')
                assert sheet.max_row < 200 and sheet.max_column <= (8 if sheet.title == 'Товары' else 7)
                # Inspect blanks in memory without writing any extra cells to the XLSX.
                for row, column in ((sheet.max_row + 2, 1), (2, sheet.max_column + 2)):
                    cell = sheet.cell(row, column)
                    assert cell.value is None and not cell.has_style
                    assert cell.fill.patternType is None
                    assert all(side.style is None for side in
                               (cell.border.left, cell.border.right, cell.border.top, cell.border.bottom))
    finally:
        book.close()


def test_plain_tables_preserve_currencies_ranking_and_summary(sample_result, tmp_path):
    result = deepcopy(sample_result)
    result['action_monthly_dynamics'] = (('2025-01', 2, 1, 1, 1), ('2025-02', 5, 2, 3, 2))
    result['order_monthly_dynamics'] = (('2025-01', 1, '40.25', 1, '50.75'), ('2025-02', 2, '80.50', 3, '152.25'))
    result['favorite_user_activity_distribution'] = (('1', 1, 100.),)
    result['top_viewed_products'] = tuple((f'{i:06}', 'Одинаковое название', 30 - i, 1, 2, 3) for i in range(20))
    result['top_favorited_products'] = result['top_viewed_products']
    path = tmp_path / 'charts.xlsx'
    export.export_statistics(result, path)
    with ZipFile(path) as archive:
        assert not any(name.startswith(('xl/charts/', 'xl/drawings/')) for name in archive.namelist())
        assert not any(name.startswith('xl/media/') for name in archive.namelist())
    book = load_workbook(path)
    try:
        assert assert_minimal_tables(book)
        summary = book['Сводка']
        for label, field in [('Количество взаимодействий', 'total_interactions'), ('Количество заказов', 'orders'),
                             ('Количество позиций в заказах', 'order_lines'), ('Профили текущего среза', 'customers'),
                             ('Клиенты с товарными взаимодействиями', 'interaction_users'), ('Покупатели', 'purchase_users'),
                             ('Повторные покупатели', 'repeat_buyers'), ('Доля повторных покупателей', 'repeat_buyer_rate')]:
            assert kpi(summary, label).value == result[field]
        assert kpi(summary, 'Доля повторных покупателей').number_format == export.PERCENT_FORMAT
        assert len([row for row in book['Товары'].values if row[0] == '000019']) == 2
        for month, _, rub, _, kzt in result['order_monthly_dynamics']:
            row = find(book['Заказы'], month[5:] + '.' + month[:4])
            assert Decimal(str(row[2].value)) == Decimal(rub)
            assert Decimal(str(row[4].value)) == Decimal(kzt)
    finally:
        book.close()


@pytest.mark.parametrize('prefix', ['=', '+', '-', '@'])
def test_all_external_strings_literal_no_links_or_pii(sample_result, tmp_path, prefix):
    label = prefix + 'HYPERLINK("https://invalid.example", "A&B <текст>")'
    sample_result['top_viewed_products'] = ((prefix + '000001', label, 2, 1, 0, 1),)
    sample_result['store_statistics'] = (('RUB', 'private-internal-key', label, 1, 1, 1, '4', '40', '40'),)
    sample_result['warnings'] = (label,)
    # Unrelated/extra metadata cannot enter this whitelist-based projection.
    sample_result['customer_id'] = 'DO-NOT-EXPORT-CUSTOMER'
    sample_result['local_path'] = 'C:/DO-NOT-EXPORT-PATH/secret'
    path = tmp_path / 'safe.xlsx'
    export.export_statistics(sample_result, path)
    book = load_workbook(path)
    try:
        assert all(cell.data_type != 'f' and cell.hyperlink is None
                   for sheet in book for row in sheet for cell in row)
        assert find(book['Товары'], prefix + '000001')[0].data_type == 's'
        assert find(book['Заказы'], label)[0].data_type == 's'
    finally:
        book.close()
    with ZipFile(path) as archive:
        assert not any('externalLink' in name or 'vba' in name.lower() for name in archive.namelist())
        text = '\n'.join(archive.read(name).decode('utf-8') for name in archive.namelist() if name.endswith('.xml'))
        assert all(secret not in text for secret in ('DO-NOT-EXPORT-CUSTOMER', 'DO-NOT-EXPORT-PATH', 'private-internal-key'))


def test_percentages_and_quantity_formats_do_not_change_values(sample_result, tmp_path):
    sample_result['active_buyer_rate'] = 43.48
    sample_result['repeat_buyer_rate'] = 12.34
    sample_result['namespaces'] = (('offline1C', 10000, 4348, 5652, 43.48),)
    path = tmp_path / 'rates.xlsx'
    export.export_statistics(sample_result, path)
    book = load_workbook(path)
    try:
        cell = find(book['Клиенты'], 'Доля активных клиентов, совершивших покупку, %')[1]
        assert cell.value == 43.48 and cell.number_format == export.PERCENT_FORMAT
        assert find(book['Товары'], 'offline1C')[4].number_format == export.PERCENT_FORMAT
        assert kpi(book['Сводка'], 'Доля повторных покупателей').value == 12.34
    finally:
        book.close()


def test_long_ranges_names_and_filter_context_remain_complete(sample_result, tmp_path):
    collections = tuple(f'{i:03} ' + 'Очень длинное наименование коллекции ' * 6 for i in range(90))
    sample_result['analysis_filter'] = AnalysisFilter(collections=collections).to_dict()
    name = 'Длинное название магазина ' * 14
    sample_result['store_statistics'] = tuple(('RUB', str(i), f'{name}{i}', 1, 1, 1, '1', '10.20', '10.20') for i in range(240))
    sample_result['action_monthly_dynamics'] = tuple((f'{2020 + i // 12}-{i % 12 + 1:02}', i + 1, 1, 0, 0) for i in range(72))
    sample_result['product_season_statistics'] = tuple((name, 1, 1, 1, 1, '1') for name in collections)
    sample_result['warnings'] = ('Важное предупреждение ' * 100,)
    path = tmp_path / 'long.xlsx'
    export.export_statistics(sample_result, path)
    book = load_workbook(path)
    try:
        assert assert_minimal_tables(book)
        assert find(book['Сводка'], 'Сезон')[1].value.startswith('Выбрано значений: 90.')
        metadata = [row[1] for row in book['Техническая информация'].values if row[0] == 'Сезон']
        assert tuple(metadata) == AnalysisFilter.from_dict(sample_result['analysis_filter']).collections
        assert kpi(book['Сводка'], 'Количество взаимодействий').row <= 20
        assert len([row for row in book['Заказы'].values if isinstance(row[0], str) and row[0].startswith(name)]) == 240
        for month, views, *_ in sample_result['action_monthly_dynamics']:
            assert find(book['Действия'], month[5:] + '.' + month[:4])[1].value == views
    finally:
        book.close()


def test_all_snapshot_fields_consumed_no_mutation(sample_result, tmp_path):
    class Observed(dict):
        used = set()

        def __getitem__(self, name):
            self.used.add(name)
            return super().__getitem__(name)
    original = deepcopy(sample_result)
    observed = Observed(sample_result)
    export.export_statistics(observed, tmp_path / 'all.xlsx')
    assert set(sample_result) - {'warnings'} <= observed.used
    assert 'warnings' not in observed.used
    assert observed == original


def test_no_export_source_reads_even_with_valid_cached_snapshot(sample_result, tmp_path, monkeypatch):
    import sqlite3
    from Application.statistics_cache import validate_result
    validate_result(sample_result)

    def forbidden(*args, **kwargs):
        raise AssertionError('Export attempted a source read')
    monkeypatch.setattr(sqlite3, 'connect', forbidden)
    monkeypatch.setattr(Path, 'read_text', forbidden)
    monkeypatch.setattr(Path, 'read_bytes', forbidden)
    export.export_statistics(sample_result, tmp_path / 'offline.xlsx')


@pytest.mark.parametrize('missing_profiles', [False, True])
def test_empty_valid_snapshot_has_only_real_tables(sample_result, tmp_path, missing_profiles):
    result = empty_snapshot(sample_result, missing_profiles=missing_profiles)
    path = tmp_path / 'empty.xlsx'
    export.export_statistics(result, path)
    book = load_workbook(path)
    try:
        assert all(not sheet._charts for sheet in book)
        assert kpi(book['Сводка'], 'Профили текущего среза').value == ('Нет данных' if missing_profiles else 0)
        assert kpi(book['Сводка'], 'Количество взаимодействий').value == 0
        assert find(book['Клиенты'], 'Средний возраст')[1].value is None
        assert assert_minimal_tables(book)
        assert all('нет данных для диаграмм' not in str(cell.value).lower() for sheet in book for row in sheet for cell in row)
    finally:
        book.close()


def test_valid_full_and_cache_json_roundtrip_use_same_snapshot(sample_result, tmp_path):
    import json
    from Application.statistics_cache import validate_result
    result = complete_snapshot(sample_result)
    cached = validate_result(json.loads(json.dumps(result)))
    for index, snapshot in enumerate((result, cached)):
        path = tmp_path / f'backend-independent-{index}.xlsx'
        export.export_statistics(snapshot, path)
        book = load_workbook(path)
        try:
            assert assert_minimal_tables(book)
            assert kpi(book['Сводка'], 'Количество взаимодействий').value == 5
            assert find(book['Заказы'], 'RUB')[4].value == 40
            assert find(book['Заказы'], 'KZT')[4].value == 50
        finally:
            book.close()


def test_stress_valid_snapshot_preserves_all_detail_rows(sample_result, tmp_path):
    result = stress_snapshot(sample_result)
    original = deepcopy(result)
    path = tmp_path / 'stress.xlsx'
    export.export_statistics(result, path)
    assert result == original
    book = load_workbook(path)
    try:
        assert assert_minimal_tables(book)
        assert len([row for row in book['Заказы'].values if 'Длинное название тестового магазина' in str(row[0])]) == 240
        for row in result['top_purchased_products']:
            assert find(book['Товары'], row[0])[1].value == row[1]
        months = [row[0][5:] + '.' + row[0][:4] for row in result['action_monthly_dynamics']]
        assert [row[0] for row in book['Действия'].values if row[0] in months] == months
    finally:
        book.close()


@pytest.mark.parametrize('amount', ['1500.25', '1500000.25', '1500000000.25'])
def test_money_values_remain_numeric_and_unscaled(sample_result, tmp_path, amount):
    sample_result['order_monthly_dynamics'] = (('2025-01', 1, amount, 0, '0'),)
    path = tmp_path / 'money.xlsx'
    export.export_statistics(sample_result, path)
    book = load_workbook(path)
    try:
        row = find(book['Заказы'], '01.2025')
        assert row[2].data_type == 'n' and row[2].number_format == '#,##0.00'
        assert Decimal(str(row[2].value)) == Decimal(amount)
        assert row[4].value == 0
        assert assert_minimal_tables(book)
    finally:
        book.close()


def test_collection_table_keeps_source_order_and_all_rows(sample_result, tmp_path):
    original = tuple((f'Коллекция {i:03}', 1, 1, 1, i % 17, '1') for i in range(120))
    sample_result['product_season_statistics'] = original
    path = tmp_path / 'collections.xlsx'
    export.export_statistics(sample_result, path)
    book = load_workbook(path)
    try:
        exported = [row[:6] for row in book['Товары'].values if str(row[0]).startswith('Коллекция ')]
        assert exported == [(*row[:5], int(row[5])) for row in original]
        assert sample_result['product_season_statistics'] == original
        assert assert_minimal_tables(book)
    finally:
        book.close()
