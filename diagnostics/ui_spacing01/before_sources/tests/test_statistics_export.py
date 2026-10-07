"""Offline XLSX snapshot export, numeric types and atomic publication."""
from copy import deepcopy
from datetime import datetime
from decimal import Decimal

from openpyxl import load_workbook
import pytest

from Application import statistics_export as export
from Application.analysis_filter import AnalysisFilter
from tabs import test_dataset_statistics_tab as samples

sample_result = samples.sample_result


def rows(sheet):
    return list(sheet.values)


def find(sheet, label):
    return next(row for row in sheet.iter_rows() if row[0].value == label)


def test_workbook_snapshot_metadata_types_and_sections(sample_result, tmp_path):
    result = sample_result
    result['analysis_filter'] = AnalysisFilter('2025-01-02', '2025-12-30', ('Рубашки', None), None).to_dict()
    result['warnings'] = ('Проверка предупреждения',)
    result['product_season_statistics'] = (('Зима 2025', 1, 2, 0, 1, '4'),)
    result['order_financials'] = (('RUB', 1, 1, '4', '3901955713.80', '3901955713.80', '3901955713.80'),
                                  ('KZT', 1, 1, '1', '123456.78', '123456.78', '123456.78'))
    result['store_statistics'] = tuple(('RUB', str(i), f'Магазин {i}', 1, 1, 1, '1', '10.20', '10.20') for i in range(25)) + (
        ('KZT', 'kz', 'Казахстан', 1, 1, 1, '1', '123456.78', '123456.78'),)
    original = deepcopy(result)
    path = tmp_path / 'statistics.xlsx'
    export.export_statistics(result, path)
    assert result == original
    book = load_workbook(path)
    assert book.sheetnames == list(export.SHEET_NAMES)
    assert all(sheet.max_row > 10 and not sheet._charts and not sheet.merged_cells for sheet in book)
    summary = book['Сводка']
    assert find(summary, 'Дата и время расчета')[1].value == datetime.fromisoformat(result['calculated_at']).astimezone().strftime('%d.%m.%Y %H:%M:%S')
    assert find(summary, 'Вид номенклатуры')[1].value == 'Рубашки; Не указано'
    assert find(summary, 'Сезон')[1].value == 'Все значения'
    assert find(summary, 'Дата начала')[1].value == datetime(2025, 1, 2)
    assert find(summary, 'Количество взаимодействий')[1].value == 3
    assert find(summary, 'Проверка предупреждения')
    orders = book['Заказы']
    for currency, expected in [('RUB', '3901955713.80'), ('KZT', '123456.78')]:
        row = find(orders, currency)
        assert row[4].data_type == 'n'
        assert Decimal(str(row[4].value)) == Decimal(expected)
        assert row[4].number_format == '#,##0.00'
    assert len([r for r in rows(orders) if isinstance(r[0], str) and r[0].startswith('Магазин ') and r[1] == 1]) == 25
    assert find(orders, 'Магазины и каналы — RUB') and find(orders, 'Магазины и каналы — KZT')
    assert find(book['Товары'], 'Зима 2025')
    assert find(book['Товары'], '000001')[0].data_type == 's'
    assert find(book['Клиенты'], 'Средний возраст')[1].value is None
    assert find(book['Клиенты'], 'Доля активных клиентов, совершивших покупку, %')[1].value == 50
    assert find(book['Техническая информация'], 'Источник')[2].value == 'Источник данных'
    assert find(book['Техническая информация'], 'Заказы вне периода статистики')[1].data_type == 'n'
    assert find(book['Техническая информация'], 'Проверка предупреждения')
    book.close()


def test_literal_labels_not_formulas_and_decimal_quantity(sample_result, tmp_path):
    sample_result['top_purchased_products'] = (('=1+1', 'Название', 1, 1, '1234.125', 2, 0),)
    path = tmp_path / 'literal.xlsx'
    export.export_statistics(sample_result, path)
    book = load_workbook(path)
    row = find(book['Товары'], '=1+1')
    assert row[0].data_type == 's'
    assert row[4].data_type == 'n' and Decimal(str(row[4].value)) == Decimal('1234.125')
    book.close()


@pytest.mark.parametrize('phase', ['save', 'replace'])
def test_atomic_failure_keeps_destination_and_cleans_temp(sample_result, tmp_path, monkeypatch, phase):
    path = tmp_path / 'existing.xlsx'
    path.write_bytes(b'previous file')
    def fail(*args):
        raise PermissionError('synthetic')
    if phase == 'save':
        monkeypatch.setattr(export.Workbook, 'save', fail)
    else:
        monkeypatch.setattr(export.os, 'replace', fail)
    with pytest.raises(PermissionError):
        export.export_statistics(sample_result, path)
    assert path.read_bytes() == b'previous file'
    assert list(tmp_path.iterdir()) == [path]


def test_no_raw_or_qt_dependency(tmp_path):
    import subprocess
    import sys
    result = subprocess.run([sys.executable, '-c',
        "import sys; import Application.statistics_export; "
        "assert 'PyQt6' not in sys.modules; "
        "assert 'Application.dataset_statistics' not in sys.modules; "
        "assert 'Application.mindbox.raw_reader' not in sys.modules"], capture_output=True, text=True)
    assert result.returncode == 0, result.stderr


def test_raw_actions_in_technical_and_filter_independent(sample_result, tmp_path):
    sample_result['source_actions'] = 100
    sample_result['action_types'] = (('ProsmotrProdukta', 90), ('ProsmotrKategoriiProduktov', 10))
    raw_rows = []
    for filtered in (False, True):
        sample_result['analysis_filter'] = AnalysisFilter(nomenclature_types=('Рубашки',) if filtered else None).to_dict()
        path = tmp_path / f'raw-{filtered}.xlsx'
        export.export_statistics(sample_result, path)
        book = load_workbook(path)
        assert find(book['Сводка'], 'Количество взаимодействий')[1].value == sample_result['total_interactions']
        business = {cell.value for row in book['Действия'] for cell in row}
        assert 'Системное название' not in business and 'Классификация событий' not in business
        technical = book['Техническая информация']
        assert find(technical, 'Исходные Actions')[1].value == 100
        raw = [find(technical, name) for name, _ in sample_result['action_types']]
        assert all(cell.data_type == 'n' for row in raw for cell in row[1:3])
        values = [(row[1].value, row[2].value) for row in raw]
        assert sum(v[1] for v in values) == pytest.approx(100)
        assert all(v[1] <= 100 for v in values)
        raw_rows.append(values)
        book.close()
    assert raw_rows[0] == raw_rows[1]
